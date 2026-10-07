"""事件窗量化对照（event_impact_panel）接线测试 — 开关/挂载点/注册表/双端/LLM 串行化。

覆盖（迭代 3 验收）：
  - 开关注册表：event_window_impact 已登记（实验组、默认关、影响报告）
  - 挂载点：seam 开关关闭零行为、开关注入契约、下游异常只告警不外抛
  - 注册表：无独立条目（并入新闻章）、号序连续、页签名/导航/两端派生点无残留
  - 关态回退：契约缺席 → 新闻章内区块零渲染；新闻父章关 → 区块随父章隐藏
  - 双端单源：财经新闻页签尾部区块单元格 ⊇ view 同源行；口径标注与事件表同现
  - 分歧例块：全部分歧例 100% 进块；无分歧/不可用 → 空串（附录零贡献）
  - 模板接线：partial 存在且被 report_template include 在新闻章内（block-title 区块级守卫）
  - LLM 编排串行化：开关开启 → 新闻先行（seam 注入后才提交 LLM）；
    开关关闭 → 并行原路径且 seam 不被调用

测试隔离：全链 mock/注入替身，不触网、不读写真实数据。
"""

from __future__ import annotations

import re
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

from src.python.config.features import feature_switch_registry, is_feature_enabled
from src.python.core.constants import PROJECT_ROOT
from src.python.core.registry import _REPORT_SECTION_DEFAULT
from src.python.report import _experimental_seams as seams
from src.python.report.event_impact_panel import (
    EVENT_TABLE_HEADER,
    LIMITATIONS_NOTE,
    build_event_impact_panel,
    build_event_impact_view,
    write_event_impact_footer,
)
from src.python.report.progress import ProgressReporter

pytestmark = [pytest.mark.unit, pytest.mark.unit_report, pytest.mark.usefixtures("offline_external_sources")]

_TMPL_DIR = Path(PROJECT_ROOT) / "src" / "static" / "tmpl"

_PANEL_BUILD = "src.python.report.event_impact_panel.build_event_impact_panel"


# 交易日历注入（与 test_event_impact_panel 同构，隔离 akshare 网络调用）
_CAL = {
    (datetime(2026, 8, 24) + timedelta(days=i)).strftime("%Y-%m-%d")
    for i in range((datetime(2026, 10, 30) - datetime(2026, 8, 24)).days + 1)
    if (datetime(2026, 8, 24) + timedelta(days=i)).weekday() < 5
}


@pytest.fixture(autouse=True)
def _calendar(offline_external_sources):
    """注入交易日历（patch 持有者模块，隔离 akshare 网络调用）。

    显式依赖 ``offline_external_sources``：offline 桩也改写同一属性，
    必须让 offline **先装配**（其 monkeypatch 保存真函数）→ 本 fixture 后装配；
    逆序终化时才能依次还原到真函数，否则 offline 会把本 fixture 的 MagicMock
    存为「原值」并在最后一步还原出来（泄漏到后续用例）。
    """
    with patch("src.python.core.trading_calendar._get_trading_calendar", return_value=set(_CAL)):
        yield


class _RecordingReporter(ProgressReporter):
    """记录各级消息的进度报告器（默认 ProgressReporter 各方法为空实现）。"""

    def __init__(self) -> None:
        super().__init__()
        self.messages: list[tuple[str, str]] = []

    def warn(self, msg: str) -> None:
        self.messages.append(("warn", msg))

    def warns(self) -> list[str]:
        return [msg for level, msg in self.messages if level == "warn"]


@pytest.fixture
def reporter() -> _RecordingReporter:
    return _RecordingReporter()


def _event(index: int, match: str, polarity: str, day: str | None = None) -> dict:
    """事件行 fixture（字段形状与 analysis.event_window_impact 输出一致）。"""
    return {
        "event_date": day or f"2026-09-{10 + index:02d}",
        "anchor_date": day or f"2026-09-{10 + index:02d}",
        "title": f"示例新闻 {index}",
        "code": "300750",
        "matched_keyword": "宁德时代",
        "polarity": polarity,
        "asset_return": -0.0202,
        "index_return": -0.0401,
        "excess_return": 0.0199,
        "direction": "下",
        "match": match,
    }


def _contract(events: list[dict], *, available: bool = True) -> dict:
    match_counts: dict[str, int] = {}
    for e in events:
        match_counts[e["match"]] = match_counts.get(e["match"], 0) + 1
    return {
        "available": available,
        "reason": "" if available else "无可用事件（新闻为空/窗口未走完/行情或基准缺失）",
        "events": events,
        "degraded": [],
        "degraded_count": 0,
        "skipped": {"bad_date": 0, "unmatched": 3, "capped": 0},
        "news_total": 10,
        "polarity_known": 8,
        "match_counts": match_counts,
        "benchmark": {"code": "sh000300", "name": "沪深300", "source": "default"},
        "window": {"before": 5, "after": 5},
        "limitations_note": LIMITATIONS_NOTE,
        "prompt_block": "",
    }


# ── 开关注册表 ───────────────────────────────────────────


class TestFeatureSwitchRegistry:
    """event_window_impact 已登记且默认关、影响报告输出。"""

    def test_switch_registered_in_experimental_group_off_by_default(self):
        sw = feature_switch_registry["event_window_impact"]
        assert sw.label == "事件窗量化对照"
        assert sw.group == "experimental"
        assert sw.default is False
        assert sw.affects_report is True
        assert is_feature_enabled("event_window_impact") is False


# ── 挂载点（seam） ──────────────────────────────────────


class TestEventImpactMount:
    """事件窗量化对照挂载点（实验组开关，产出 ``event_impact_data`` 契约）。"""

    def test_flag_off_is_inert(self, reporter):
        """开关关闭 → 不装配、不注入、不告警（零导入零成本）。"""
        pipeline_data: dict = {}
        with patch("src.python.config.features.is_feature_enabled", return_value=False):
            seams.inject_event_impact_data(pipeline_data, reporter, news_data=[], holdings=[])
        assert pipeline_data == {}
        assert reporter.warns() == []

    def test_flag_on_injects_contract(self, reporter):
        """开关打开 → 契约写入 pipeline_data，新闻/持仓/宽基池透传。"""
        pipeline_data: dict = {}
        news = [{"title": "x"}]
        with (
            patch("src.python.config.features.is_feature_enabled", return_value=True),
            patch(_PANEL_BUILD, return_value={"available": True}) as build,
        ):
            seams.inject_event_impact_data(
                pipeline_data,
                reporter,
                news_data=news,
                holdings=["H"],
                penetrated_assets=[{"name": "宁德时代"}],
                comparison_indices={"sh000300": "沪深300"},
            )
        assert pipeline_data["event_impact_data"] == {"available": True}
        assert build.call_args.args[0] == ["H"]
        assert build.call_args.args[1] is news
        assert build.call_args.kwargs["comparison_indices"] == {"sh000300": "沪深300"}

    def test_none_pipeline_data_does_not_raise(self, reporter):
        """pipeline_data=None（干跑）不得因注入而抛异常。"""
        with (
            patch("src.python.config.features.is_feature_enabled", return_value=True),
            patch(_PANEL_BUILD, return_value={"available": True}),
        ):
            seams.inject_event_impact_data(None, reporter, news_data=[], holdings=[])
        assert reporter.warns() == []

    def test_downstream_exception_warns_only(self, reporter):
        """下游异常 → 一条告警 + 契约缺席，不外抛、不中断报告。"""
        pipeline_data: dict = {}
        with (
            patch("src.python.config.features.is_feature_enabled", return_value=True),
            patch(_PANEL_BUILD, side_effect=RuntimeError("boom")),
        ):
            seams.inject_event_impact_data(pipeline_data, reporter, news_data=[], holdings=[])
        assert pipeline_data == {}
        assert len(reporter.warns()) == 1
        assert "事件窗量化对照" in reporter.warns()[0]


# ── 注册表 / 导航 / 页签名 ──────────────────────────────


class TestRegistryWiring:
    """事件窗并入新闻章：注册表无独立条目，两端派生点（号序/页签名/导航）无残留。"""

    def test_no_section_entry(self):
        assert all(s["key"] != "event_impact" for s in _REPORT_SECTION_DEFAULT)

    def test_numbering_contiguous_after_removal(self):
        """移除独立条目后号序仍连续（区段无断号）。"""
        keys = [s["key"] for s in _REPORT_SECTION_DEFAULT]
        numbers = [s["number"] for s in _REPORT_SECTION_DEFAULT]
        assert "event_impact" not in keys
        assert numbers == list(range(1, len(keys) + 1))

    def test_no_sheet_name_registered(self):
        from src.python.core.registry import _REPORT_SHEET_NAMES

        assert "event_impact" not in _REPORT_SHEET_NAMES

    def test_no_nav_group_entry(self):
        from src.python.core.registry import get_report_section_keys

        assert "event_impact" not in get_report_section_keys()

    def test_no_board_or_data_flag_on_either_side(self):
        """两端可见性派生点均无事件窗键（区块随父章，不参与章级判定）。"""
        from src.python.report.excel_sheet_factory import build_data_availability
        from src.python.report.html_writer_nav import _compute_section_visibility

        _nums, visible, _ = _compute_section_visibility(
            list(_REPORT_SECTION_DEFAULT), {}, {}, {}, {}, include_news=True, llm_enabled_flag=True
        )
        assert "event_impact" not in visible
        assert "event_impact_data" not in build_data_availability()


# ── 关态回退：区块不占编号，随父章可见性 ───────────────


class TestOffStateVisibility:
    """区块不占章级可见性/编号；新闻父章关 → 区块随父章一并隐藏。"""

    def test_news_numbering_contiguous_without_event_key(self):
        from src.python.core.registry import get_report_section_order
        from src.python.report.html_writer_nav import _compute_section_visibility

        order = get_report_section_order()
        visible_numbers, visible_dict, _ = _compute_section_visibility(
            order, {}, {}, {}, {}, include_news=True, llm_enabled_flag=True
        )
        assert visible_dict["news_correlation"] is True
        assert "event_impact" not in visible_dict
        assert list(visible_numbers.values()) == list(range(1, len(visible_numbers) + 1))

    def test_news_board_off_hides_news_parent_of_block(self):
        """新闻 board 层关闭 → 父章隐藏，事件窗区块随之不可达（从属关系）。"""
        from src.python.core.registry import get_report_section_order
        from src.python.report.html_writer_nav import _compute_section_visibility

        order = get_report_section_order()
        visible_numbers, visible_dict, _ = _compute_section_visibility(
            order, {}, {}, {}, {}, include_news=True, llm_enabled_flag=True, enable_news=False
        )
        assert visible_dict["news_correlation"] is False
        assert "news_correlation" not in visible_numbers

    def test_build_data_availability_has_no_event_key(self):
        from src.python.report.excel_sheet_factory import build_data_availability

        assert "event_impact_data" not in build_data_availability()


# ── 双端单源：view ↔ 财经新闻页签尾部区块 ──────────────


class TestViewSheetSingleSource:
    """view 预格式化行与财经新闻页签尾部区块逐字节同文；口径标注与事件表同现。"""

    def test_footer_rows_superset_of_view_rows(self):
        from openpyxl import Workbook

        contract = _contract([_event(1, "一致", "利空"), _event(2, "分歧", "利好")])
        view = build_event_impact_view(contract)

        wb = Workbook()
        ws = wb.active
        write_event_impact_footer(ws, contract, start_row=5)
        assert ws.cell(row=5, column=1).value == "事件窗量化对照", "区块标题必须落在 start_row"
        cells = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c is not None]

        for line in view["summary_lines"]:
            assert line in cells, "口径摘要行未逐字节写入页签尾部区块"
        for row in view["event_rows"]:
            for cell in row:
                assert cell in cells, f"事件行单元格缺失: {cell}"
        assert view["limitations_note"] in cells, "口径标注必须与事件表同现（双端成对）"
        assert LIMITATIONS_NOTE in cells

    def test_event_header_matches_table_constant(self):
        view = build_event_impact_view(_contract([_event(1, "一致", "利空")]))
        assert view["event_header"] == list(EVENT_TABLE_HEADER)
        assert view["event_rows"][0][9] in {"一致", "分歧", "中性不判"}  # 比对列在末位

    def test_placeholder_when_unavailable(self):
        from openpyxl import Workbook

        contract = _contract([], available=False)
        view = build_event_impact_view(contract)
        assert view["available"] is False
        assert view["reason"] == contract["reason"]

        wb = Workbook()
        ws = wb.active
        write_event_impact_footer(ws, contract, start_row=3)
        texts = [str(c) for row in ws.iter_rows(values_only=True) for c in row if c is not None]
        assert contract["reason"] in texts, "不可用时页签尾部区块必须写占位原因"


# ── 分歧例块：召回 100% / 空串回退 ──────────────────────


class TestPromptBlock:
    """prompt_block 只收分歧例：全部分歧例进块、无分歧为空、不可用为空。"""

    @staticmethod
    def _block(events: list[dict], *, available: bool = True) -> str:
        from src.python.report.event_impact_panel import _build_prompt_block

        contract = _contract(events, available=available)
        view = build_event_impact_view(contract)
        return _build_prompt_block(view)

    def test_all_divergent_events_reach_block(self):
        events = [_event(1, "分歧", "利好"), _event(2, "分歧", "利好"), _event(3, "分歧", "利好", day="2026-09-25")]
        block = self._block(events)
        assert block, "存在分歧例时块必须非空"
        for e in events:
            assert e["event_date"] in block, f"分歧例未进块（召回率不足 100%）: {e['event_date']}"
            assert e["code"] in block

    def test_non_divergent_rows_excluded(self):
        events = [_event(1, "一致", "利空"), _event(2, "中性不判", "中性"), _event(3, "分歧", "利好")]
        block = self._block(events)
        assert events[2]["event_date"] in block
        assert events[0]["event_date"] not in block, "一致例不应进分歧例块"
        assert events[1]["event_date"] not in block, "中性不判不应进分歧例块"

    def test_empty_without_divergence(self):
        assert self._block([_event(1, "一致", "利空")]) == "", "无分歧时块必须为空（附录零贡献）"

    def test_empty_when_unavailable(self):
        assert self._block([], available=False) == "", "契约不可用时块必须为空"


# ── 模板接线（源码级） ───────────────────────────────────


class TestTemplateWiring:
    """partial 已建且被主模板 include 在新闻章内，按 event_impact_view 区块级守卫。"""

    def test_partial_exists_and_guarded(self):
        partial = (_TMPL_DIR / "partials" / "event_impact_section.html").read_text(encoding="utf-8")
        assert "{% if event_impact_view %}" in partial
        assert 'section_visible("event_impact")' not in partial  # 不再有章级可见性键
        assert "event_impact_view" in partial
        assert "limitations_note" in partial  # 标注经同源 view 渲染
        assert "sec-event_impact" in partial
        assert 'class="block-title"' in partial  # 章内区块标题级（与持仓基本面同构）

    def test_main_template_includes_partial_inside_news_chapter(self):
        """event_impact include 位于新闻章 partial 内，主模板经新闻章 include 挂载，全库仅一处。"""
        tmpl = (_TMPL_DIR / "report_template.html").read_text(encoding="utf-8")
        news = (_TMPL_DIR / "partials" / "news_correlation_section.html").read_text(encoding="utf-8")
        assert '{% include "partials/news_correlation_section.html" with context %}' in tmpl
        assert '{% include "partials/event_impact_section.html" with context %}' in news
        # include 在新闻章起点之后，且新闻章 partial 不含下一章（章边界 = partial 边界）
        assert news.index('id="sec-news_correlation"') < news.index("partials/event_impact_section.html")
        assert 'id="sec-global_macro"' not in news
        sources = [tmpl] + [p.read_text(encoding="utf-8") for p in sorted((_TMPL_DIR / "partials").glob("*.html"))]
        assert sum(s.count("partials/event_impact_section.html") for s in sources) == 1


# ── LLM 编排串行化（_fetch_llm_and_news） ───────────────


class TestNewsLlmSerialization:
    """开关开启 → 新闻先行（seam 注入后提交 LLM）；关闭 → 并行原路径且不注入。"""

    @staticmethod
    def _completed(result):
        from concurrent.futures import Future

        fut = Future()
        fut.set_result(result)
        return fut

    @staticmethod
    def _invoke(flag: bool, record: list[str]):
        with (
            patch("src.python.config.features.is_feature_enabled", return_value=flag),
            patch(
                "src.python.report._llm_news._submit_news_future",
                side_effect=lambda *a, **k: (
                    record.append("submit_news") or TestNewsLlmSerialization._completed(([], {}, False))
                ),
            ),
            patch(
                "src.python.report._llm_news._submit_llm_future",
                side_effect=lambda *a, **k: (
                    record.append("submit_llm")
                    or TestNewsLlmSerialization._completed((None, None, None, None, False, False, False, False))
                ),
            ),
            patch(
                "src.python.report._experimental_seams.inject_event_impact_data",
                side_effect=lambda *a, **k: record.append("inject"),
            ),
        ):
            from src.python.report._llm_news import _fetch_llm_and_news

            _fetch_llm_and_news(
                [],
                {"news_top_count": 10, "penetrated_assets": None},
                None,
                False,  # force
                {"purchase_status_data": {}},
                True,  # enable_news
                True,  # enable_llm
                MagicMock(),
                comparison_indices={"sh000300": "沪深300"},
            )

    def test_flag_on_serializes_news_before_llm(self):
        record: list[str] = []
        self._invoke(True, record)
        assert record.index("submit_news") < record.index("inject") < record.index("submit_llm"), (
            f"开关开启时应为 新闻 → seam 注入 → LLM 的串行序，实际 {record}"
        )

    def test_flag_off_keeps_parallel_and_skips_seam(self):
        record: list[str] = []
        self._invoke(False, record)
        assert "inject" not in record, "开关关闭时 seam 不得被调用"
        assert record.count("submit_news") == 1 and record.count("submit_llm") == 1, (
            f"开关关闭走并行原路径（双提交），实际 {record}"
        )


# ── 契约装配（注入替身，全离线） ─────────────────────────


class TestPanelAssembly:
    """build_event_impact_panel：注入取数 → 事件表 + 分歧例块落契约。"""

    def test_assembles_contract_with_block(self):
        class _H:
            name = "宁德时代"
            code = "300750"
            shares = 100.0
            cost_price = 200.0

        news = [
            {
                "title": "宁德时代利好新闻",
                "ctime": "2026-09-10 09:30:00",
                "matched_keywords": ["宁德时代"],
                "llm_analysis": "[高][利好] 关联判定",
            }
        ]
        bars = [{"date": f"2026-09-{d:02d}", "close": 100.0 + d} for d in range(1, 30)]
        index_bars = [{"date": f"2026-09-{d:02d}", "close": 3000.0 + d} for d in range(1, 30)]

        contract = build_event_impact_panel(
            [_H()],
            news,
            None,
            benchmark_text_getter=lambda _c: (_ for _ in ()).throw(RuntimeError("无文本")),
            fetch_bars=lambda _c: bars,
            fetch_index=lambda _c: index_bars,
        )
        assert contract["available"] is True
        assert contract["news_total"] == 1
        assert contract["polarity_known"] == 1
        assert contract["match_counts"], "应产出比对计数"
        assert contract["window"] == {"before": 5, "after": 5}
        # 分歧例（利好 ↓）→ 块非空；一致则块为空——两种形态均由判定决定，此处只锁结构
        assert isinstance(contract["prompt_block"], str)

    def test_empty_news_yields_unavailable_contract(self):
        contract = build_event_impact_panel([], [], None, fetch_bars=lambda _c: [], fetch_index=lambda _c: [])
        assert contract["available"] is False
        assert contract["reason"]
        assert contract["prompt_block"] == ""
        view = build_event_impact_view(contract)
        assert view["available"] is False
        assert re.search(r"[一-鿿]", view["reason"]), "占位原因必须是中文文案"
        assert view["hint_line"]


class TestCalendarFixtureStackOrder:
    """日历注入与 offline 桩的装配栈序（防 patch 泄漏到后续用例）。

    缺陷回归：``_calendar`` 曾为独立 autouse（先于 ``offline_external_sources``
    装配）——offline 的 monkeypatch 会把日历 MagicMock 存为「原值」，逆序终化时
    把 MagicMock 还原出来泄漏给后续用例（实测把 ``test_market_value`` 的并发
    串行化用例打红）。修复 = ``_calendar`` 显式依赖 offline（强制 offline 先装配）。
    """

    def test_calendar_fixture_declares_offline_dependency(self):
        import inspect

        params = list(inspect.signature(_calendar).parameters)
        assert "offline_external_sources" in params, (
            "_calendar 必须显式依赖 offline_external_sources——否则 autouse 先于 offline 装配，"
            "offline 桩会把日历 MagicMock 存为原值并在终化时泄漏"
        )

    def test_stacked_teardown_restores_lower_layer(self):
        """栈序模型：后装配者先终化，逐层还原到装配前的下层对象。"""
        import src.python.core.trading_calendar as tc

        lower = tc._get_trading_calendar  # 当前栈下层（本模块 _calendar 活跃期）
        mp = pytest.MonkeyPatch()
        try:
            mp.setattr(tc, "_get_trading_calendar", lambda: set())  # 模拟 offline 后装配
            with patch("src.python.core.trading_calendar._get_trading_calendar", return_value=set(_CAL)):
                assert tc._get_trading_calendar() == set(_CAL)
            assert tc._get_trading_calendar() == set()  # 日历 patch 终化 → 还原 offline 桩
        finally:
            mp.undo()
        assert tc._get_trading_calendar is lower  # offline 终化 → 还原到栈下层（不浅栈、不越栈）
