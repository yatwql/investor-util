"""Excel 端导航三件套测试 — 页签分组配色 / 汇总页章节导航区 / 返回汇总链接。

与 HTML 端目录导航（core/registry.py 的 nav_group 单源派生）双端对账：
  - 页签 tabColor 按 NAV_GROUPS 分组上色（六组六色，键集双向一致）
  - 汇总页「章节导航」区超链接集 = 可见页签集（含顺序）
  - 各页签标题行尾「↩ 返回汇总」内部超链接
  - TOC 分组 ↔ 页签颜色同组同色（双端一致性）

运行：
  pytest src/test/unit/report/test_excel_navigation.py -v
"""

from __future__ import annotations

import re

import pytest

from src.python.core.registry import NAV_GROUPS, _REPORT_SECTION_DEFAULT, get_report_section_order

pytestmark = [pytest.mark.unit, pytest.mark.unit_report, pytest.mark.usefixtures("offline_external_sources")]

_LINK_RE = re.compile(r'=HYPERLINK\("#\'(.+?)\'!A1","')


def _make_wb():
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _all_visible_sheets():
    """用真实注册表创建全可见页签（board 全开 + data_flag_any 契约就绪）。"""
    from src.python.report.excel_sheet_factory import create_sheets

    wb = _make_wb()
    sheets = create_sheets(
        wb,
        get_report_section_order(),
        enable_fund_deep_analysis=True,
        enable_news=True,
        enable_history=True,
        enable_portfolio_evolution=True,
        enable_fundamental_snapshot=True,
        enable_action=True,
        enable_llm=True,
        data_availability={
            "position_relationship_data": True,
            "concentration_data": True,
            "financial_indicator_data": True,
            "financial_report_digest_data": True,
        },
    )
    return sheets


class TestTabColorRegistry:
    """页签配色表与 NAV_GROUPS 分组注册表的一致性。"""

    def test_color_keys_match_nav_groups_bidirectionally(self):
        """配色键集与 NAV_GROUPS 键集双向一致（无孤儿组、无未登记组）。"""
        from src.python.report.excel_sheet_factory import _GROUP_TAB_COLORS

        assert set(_GROUP_TAB_COLORS) == {g["key"] for g in NAV_GROUPS}

    def test_colors_are_unique_valid_hex(self):
        """六组六色：均为 6 位十六进制且互不相同（同色即分组不可辨）。"""
        from src.python.report.excel_sheet_factory import _GROUP_TAB_COLORS

        colors = list(_GROUP_TAB_COLORS.values())
        assert len(colors) == len(set(colors))
        for key, color in _GROUP_TAB_COLORS.items():
            assert re.fullmatch(r"[0-9A-Fa-f]{6}", color), f"{key} 配色 {color!r} 非 6 位 hex"

    def test_every_section_maps_to_a_colored_group(self):
        """每个章节条目的 nav_group 都能在配色表中找到颜色（注册表 ⊆ 配色表）。"""
        from src.python.report.excel_sheet_factory import _GROUP_TAB_COLORS

        for sec in _REPORT_SECTION_DEFAULT:
            assert sec["nav_group"] in _GROUP_TAB_COLORS, f"{sec['key']} 分组 {sec['nav_group']!r} 无配色"


class TestSheetTabColors:
    """create_sheets 创建的页签按分组上色。"""

    def test_every_created_sheet_carries_group_color(self):
        """已创建页签的 tabColor 与该章 nav_group 的组色一致（同章同组必同色）。"""
        from src.python.report.excel_sheet_factory import _GROUP_TAB_COLORS

        sheets = _all_visible_sheets()
        by_key = {sec["key"]: sec for sec in _REPORT_SECTION_DEFAULT}
        assert set(sheets) == set(by_key), "全可见场景应创建全部章节页签"
        for key, ws in sheets.items():
            color = ws.sheet_properties.tabColor
            assert color is not None, f"{key} 页签未上色"
            expected = _GROUP_TAB_COLORS[by_key[key]["nav_group"]]
            assert color.rgb.upper().endswith(expected), (
                f"{key} 页签色 {color.rgb} 与分组 {by_key[key]['nav_group']} 色 {expected} 不符"
            )


class TestSummarySectionNav:
    """汇总页「章节导航」区 — 超链接集与可见页签集双端对账。"""

    def _render_summary_nav(self, sheets):
        from unittest.mock import patch

        from src.python.report.summary import write_summary_sheet

        ws = sheets["summary"]
        titles = [w.title for w in sheets.values() if w is not None]
        with patch("src.python.report.summary.get_last_trading_day", return_value="2026-10-06"):
            write_summary_sheet(ws, 0.0, 0.0, 0.0, 0.0, section_nav=titles)
        links = []
        for row in ws.iter_rows(min_row=1, max_row=10, max_col=20):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=HYPERLINK"):
                    m = _LINK_RE.match(cell.value)
                    assert m, f"导航超链接格式异常: {cell.value!r}"
                    links.append(m.group(1))
        return links

    def test_nav_links_equal_visible_sheet_titles_in_order(self):
        """导航链接列表 = 可见页签标题列表（集合与顺序双一致，含 llm_usage 末位）。"""
        sheets = _all_visible_sheets()
        links = self._render_summary_nav(sheets)
        titles = [w.title for w in sheets.values() if w is not None]
        assert links == titles

    def test_nav_label_and_header_after_nav(self):
        """导航区带「【章节导航】」标签行，表头行紧随导航区之后（结构可发现）。"""
        sheets = _all_visible_sheets()
        ws = sheets["summary"]
        titles = [w.title for w in sheets.values() if w is not None]
        from unittest.mock import patch

        from src.python.report.summary import write_summary_sheet

        with patch("src.python.report.summary.get_last_trading_day", return_value="2026-10-06"):
            write_summary_sheet(ws, 0.0, 0.0, 0.0, 0.0, section_nav=titles)
        assert ws.cell(row=2, column=1).value == "【章节导航】"
        # 表头行 = 标题1 + 标签1 + 导航占行（向上取整）
        header_row = 2 + 1 + -(-len(titles) // 8)
        assert ws.cell(row=header_row, column=1).value == "指标"
        assert ws.freeze_panes == f"A{header_row + 1}"

    def test_no_nav_block_when_section_nav_absent(self):
        """不传 section_nav 时无导航区，表头保持在第 2 行（既有输出兼容）。"""
        from unittest.mock import patch

        from src.python.report.summary import write_summary_sheet

        from openpyxl import Workbook

        ws = Workbook().create_sheet()
        with patch("src.python.report.summary.get_last_trading_day", return_value="2026-10-06"):
            write_summary_sheet(ws, 0.0, 0.0, 0.0, 0.0)
        assert ws.cell(row=2, column=1).value == "指标"
        assert ws.cell(row=3, column=1).value != "【章节导航】"


class TestBackToSummaryLinks:
    """各页签标题行尾「↩ 返回汇总」内部超链接。"""

    def test_every_non_summary_sheet_stamped(self):
        """除汇总页外每个页签行 1 尾列有返回汇总超链接，目标为汇总页标题。"""
        from src.python.report.excel_sheet_factory import stamp_back_to_summary

        sheets = _all_visible_sheets()
        summary_title = sheets["summary"].title
        stamp_back_to_summary(sheets)
        for key, ws in sheets.items():
            value = ws.cell(row=1, column=ws.max_column).value
            if key == "summary":
                assert not (isinstance(value, str) and "返回汇总" in value), "汇总页自身不打返回标"
                continue
            assert isinstance(value, str) and value.startswith("=HYPERLINK"), f"{key} 缺少返回汇总链接"
            assert summary_title in value, f"{key} 返回链接未指向汇总页 {summary_title!r}"
            assert "返回汇总" in value

    def test_link_sits_outside_title_merge(self):
        """链接写在标题行合并区之外（不破坏标题合并且不占新行）。"""
        from src.python.report.excel_sheet_factory import stamp_back_to_summary
        from src.python.report.excel_writer import write_title_row

        sheets = _all_visible_sheets()
        ws = sheets["holdings_detail"]
        write_title_row(ws, 1, "持仓明细与分类", 6)  # 合并 A1:F1
        stamp_back_to_summary(sheets)
        assert ws.cell(row=1, column=7).value.startswith("=HYPERLINK")
        assert ws.cell(row=2, column=1).value is None, "不应新增行"


class TestDoubleEndConsistency:
    """HTML 目录分组 ↔ Excel 页签颜色（双端同源 nav_group 的接线验证）。"""

    def test_toc_group_and_tab_color_agree_per_section(self):
        """每个章节：HTML TOC 所属分组的组色 = Excel 页签 tabColor（同组同色）。"""
        from src.python.report.excel_sheet_factory import _GROUP_TAB_COLORS, create_sheets
        from src.python.report.html_writer import _build_section_nav_groups

        order = get_report_section_order()
        numbers = {s["key"]: s["number"] for s in order}
        groups = _build_section_nav_groups(order, lambda k: True, numbers)
        toc_group = {s["key"]: g["key"] for g in groups for s in g["sections"]}

        wb = _make_wb()
        sheets = create_sheets(
            wb,
            order,
            enable_fund_deep_analysis=True,
            enable_news=True,
            enable_history=True,
            enable_portfolio_evolution=True,
            enable_fundamental_snapshot=True,
            enable_action=True,
            enable_llm=True,
            data_availability={
                "position_relationship_data": True,
                "concentration_data": True,
                "financial_indicator_data": True,
                "financial_report_digest_data": True,
            },
        )
        assert set(toc_group) == set(sheets), "TOC 可见章节集应与页签集一致"
        for key, ws in sheets.items():
            expected = _GROUP_TAB_COLORS[toc_group[key]]
            assert ws.sheet_properties.tabColor.rgb.upper().endswith(expected), (
                f"{key}: TOC 分组 {toc_group[key]} 与页签色不一致"
            )
