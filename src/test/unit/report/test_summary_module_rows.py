"""汇总模块单元测试 — 模块数据行与 LLM 用量深度行写入。

测试目标：
  - _write_module_data_rows — 模块数据行写入
  - _init_llm_usage_sheet 深度行 — 深度剖析行渲染

运行：
  pytest src/test/unit/report/test_summary_module_rows.py -v
"""

from __future__ import annotations

import unittest

import pytest

pytestmark = [pytest.mark.unit, pytest.mark.unit_report]


if __name__ == "__main__":
    unittest.main()


# ═══════════════════════════════════════════════════════════
#  _write_module_data_rows — Excel 明细行单元格渲染
# ═══════════════════════════════════════════════════════════


class TestWriteModuleDataRows(unittest.TestCase):
    """测试 _write_module_data_rows 的 Excel 单元格写入逻辑。"""

    def setUp(self):
        import openpyxl

        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.start_row = 5

    def _run(self, module_info):
        """执行 _write_module_data_rows 并返回写入的单元格值字典。"""
        from src.python.report.summary_llm_usage import _write_module_data_rows

        end_row = _write_module_data_rows(self.ws, self.start_row, module_info)
        result = {}
        for r in range(self.start_row, end_row):
            row_data = {}
            for c in range(1, 11):
                cell = self.ws.cell(row=r, column=c)
                row_data[c] = cell.value
            result[r] = row_data
        return result, end_row

    def test_cache_hit_row(self):
        """缓存命中 → 蓝字'缓存'、费用'已计入原调用'、缓存✓、Thinking—。"""
        rows, end = self._run(
            [
                {
                    "key": "global_macro",
                    "name": "全球政经局势",
                    "status": "cached",
                    "status_label": "缓存",
                    "model": "deepseek-v4-flash",
                    "input_tokens": 0,
                    "output_tokens": 0,
                    "total_tokens": 0,
                    "cache_hit_tokens": 500,
                    "cost": 0.0,
                    "cached": True,
                    "thinking": False,
                    "endpoint": "",
                },
            ]
        )
        self.assertIn(self.start_row, rows)
        r = rows[self.start_row]
        self.assertEqual(r[1], "全球政经局势")
        self.assertEqual(r[2], "缓存")
        self.assertEqual(r[3], "deepseek-v4-flash")
        self.assertEqual(r[8], "已计入原调用")
        self.assertEqual(r[9], "✓")
        self.assertEqual(r[10], "—")

    def test_success_with_thinking_row(self):
        """真实调用+Thinking → 绿字'成功'、费用¥、缓存—、Thinking✓。"""
        rows, end = self._run(
            [
                {
                    "key": "expert_review",
                    "name": "智囊团深度复盘",
                    "status": "success",
                    "status_label": "成功",
                    "model": "claude-sonnet-4",
                    "input_tokens": 1500,
                    "output_tokens": 800,
                    "total_tokens": 2300,
                    "cache_hit_tokens": 0,
                    "cost": 0.005,
                    "cached": False,
                    "thinking": True,
                    "endpoint": "",
                },
            ]
        )
        r = rows[self.start_row]
        self.assertEqual(r[1], "智囊团深度复盘")
        self.assertEqual(r[2], "成功")
        self.assertEqual(r[3], "claude-sonnet-4")
        self.assertEqual(r[4], "2,300")  # total_tokens 格式化
        self.assertEqual(r[5], "1,500")
        self.assertEqual(r[6], "800")
        self.assertEqual(r[7], "—")  # cache_hit_tokens=0 → —
        self.assertIsInstance(r[8], str)
        self.assertIn("¥", str(r[8]))
        self.assertEqual(r[9], "—")
        self.assertEqual(r[10], "✓")

    def test_disabled_row(self):
        """禁用 → 灰字、模型—、费用—、缓存—、Thinking—。"""
        rows, end = self._run(
            [
                {
                    "key": "health_check",
                    "name": "持仓体检报告",
                    "status": "disabled",
                    "status_label": "已禁用",
                    "model": "",
                    "input_tokens": 0,
                    "output_tokens": 0,
                    "total_tokens": 0,
                    "cache_hit_tokens": 0,
                    "cost": 0.0,
                    "cached": False,
                    "thinking": False,
                    "endpoint": "",
                },
            ]
        )
        r = rows[self.start_row]
        self.assertEqual(r[1], "持仓体检报告")
        self.assertEqual(r[2], "已禁用")
        self.assertEqual(r[3], "—")
        self.assertEqual(r[8], "—")
        self.assertEqual(r[9], "—")
        self.assertEqual(r[10], "—")

    def test_failed_row(self):
        """失败 → 红字错误原因。"""
        rows, end = self._run(
            [
                {
                    "key": "penetration_deep",
                    "name": "穿透深度分析",
                    "status": "failed",
                    "status_label": "LLM API 调用失败",
                    "model": "",
                    "input_tokens": 0,
                    "output_tokens": 0,
                    "total_tokens": 0,
                    "cache_hit_tokens": 0,
                    "cost": 0.0,
                    "cached": False,
                    "thinking": False,
                    "endpoint": "",
                },
            ]
        )
        r = rows[self.start_row]
        self.assertEqual(r[2], "LLM API 调用失败")
        self.assertEqual(r[3], "—")

    def test_no_status_label_skipped(self):
        """status_label 为空 → 跳过该行，不写入。"""
        rows, end = self._run(
            [
                {
                    "key": "unknown",
                    "name": "未知模块",
                    "status": "unknown",
                    "status_label": "",
                    "model": "",
                    "input_tokens": 0,
                    "output_tokens": 0,
                    "total_tokens": 0,
                    "cache_hit_tokens": 0,
                    "cost": 0.0,
                    "cached": False,
                    "thinking": False,
                    "endpoint": "",
                },
            ]
        )
        self.assertEqual(end, self.start_row)
        self.assertNotIn(self.start_row, rows)

    def test_mixed_rows(self):
        """4 种状态混合 → 各行正确渲染，行号递增。"""
        rows, end = self._run(
            [
                {
                    "key": "gm",
                    "name": "全球政经局势",
                    "status": "cached",
                    "status_label": "缓存",
                    "model": "ds",
                    "cached": True,
                    "input_tokens": 0,
                    "output_tokens": 0,
                    "total_tokens": 0,
                    "cache_hit_tokens": 500,
                    "cost": 0.0,
                    "thinking": False,
                    "endpoint": "",
                },
                {
                    "key": "er",
                    "name": "智囊团深度复盘",
                    "status": "success",
                    "status_label": "成功",
                    "model": "claude",
                    "cached": False,
                    "input_tokens": 1000,
                    "output_tokens": 500,
                    "total_tokens": 1500,
                    "cache_hit_tokens": 0,
                    "cost": 0.003,
                    "thinking": True,
                    "endpoint": "",
                },
                {
                    "key": "hc",
                    "name": "持仓体检报告",
                    "status": "disabled",
                    "status_label": "已禁用",
                    "model": "",
                    "cached": False,
                    "input_tokens": 0,
                    "output_tokens": 0,
                    "total_tokens": 0,
                    "cache_hit_tokens": 0,
                    "cost": 0.0,
                    "thinking": False,
                    "endpoint": "",
                },
                {
                    "key": "pd",
                    "name": "穿透深度分析",
                    "status": "failed",
                    "status_label": "LLM API 调用失败",
                    "model": "",
                    "cached": False,
                    "input_tokens": 0,
                    "output_tokens": 0,
                    "total_tokens": 0,
                    "cache_hit_tokens": 0,
                    "cost": 0.0,
                    "thinking": False,
                    "endpoint": "",
                },
            ]
        )
        self.assertEqual(end, self.start_row + 4)
        self.assertEqual(rows[self.start_row][1], "全球政经局势")
        self.assertEqual(rows[self.start_row + 1][1], "智囊团深度复盘")
        self.assertEqual(rows[self.start_row + 2][1], "持仓体检报告")
        self.assertEqual(rows[self.start_row + 3][1], "穿透深度分析")
        # 缓存行费用
        self.assertEqual(rows[self.start_row][8], "已计入原调用")
        # 成功行费用（带 ¥）
        self.assertIn("¥", str(rows[self.start_row + 1][8]))
        # 成功行 Thinking
        self.assertEqual(rows[self.start_row + 1][10], "✓")


# ═══════════════════════════════════════════════════════════
#  _init_llm_usage_sheet — 非默认深度档位自述
# ═══════════════════════════════════════════════════════════


class TestInitLlmUsageSheetDepthLine(unittest.TestCase):
    """LLM 用量页签顶部：非默认深度档位须自述，默认档保持零噪声。"""

    def setUp(self):
        import openpyxl

        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active

    def _sheet_texts(self) -> list[str]:
        return [str(c.value) for row in self.ws.iter_rows() for c in row if c.value]

    def test_default_depth_writes_no_depth_line(self):
        from src.python.report.summary_llm_usage import _init_llm_usage_sheet

        _init_llm_usage_sheet(self.ws)
        assert not [t for t in self._sheet_texts() if "深度档位" in t]

    def test_non_default_depth_line_is_written(self):
        """非默认档位（简版）→ 页签出现含档位显示名的自述行。"""
        from unittest.mock import patch

        from src.python.llm.depth_profile import DEPTH_PROFILES
        from src.python.report.summary_llm_usage import _init_llm_usage_sheet

        profile = DEPTH_PROFILES["brief"]
        # 落点内为函数内延迟导入，故 patch 源模块属性（调用时重新绑定）
        with patch("src.python.llm.depth_profile.resolve_depth_profile", return_value=profile):
            _init_llm_usage_sheet(self.ws)

        lines = [t for t in self._sheet_texts() if "深度档位" in t]
        assert len(lines) == 1
        assert profile.label in lines[0]


if __name__ == "__main__":
    unittest.main()
