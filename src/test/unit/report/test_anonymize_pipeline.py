"""报告管线持仓匿名化接线测试 — 装配边界 pair 匿名 / 渲染层折叠 / 产物清扫。

覆盖场景（config anonymization.mode → 报告管线消费）：
  - apply_report_anonymization：off 恒等、code_display/full 两容器同代号、
    summary 折叠字典而行原样（渲染层拦截）
  - fold_detail_rows_summary：账户组内按大类折叠、合计不变、DetailRow 同构
  - prepare_report_data 装配边界注入（明细 + 约束块文本同映射脱敏）
  - Excel basic 路径内部生成门（resolve_market_data）
  - HTML 内部生成门（_render_market_value_section）+ 明细分组渲染拦截
  - mask_workbook_text：字符串单元格清扫、数值单元格不动
  - _mask_news_display_fields：新闻关键词显示字段脱敏

运行：
  pytest src/test/unit/report/test_anonymize_pipeline.py -v
"""

from __future__ import annotations

from unittest.mock import MagicMock, patch

import pytest

from src.python.core.models import Holding
from src.python.report.market_value import DetailRow

pytestmark = [
    pytest.mark.unit,
    pytest.mark.unit_report,
    pytest.mark.usefixtures("offline_external_sources"),
]

_SAMPLE_HOLDINGS = [
    Holding(account="证券", name="长江电力", code="600900", shares=100, cost_price=10.0),
    Holding(account="证券", name="贵州茅台", code="600519", shares=50, cost_price=200.0),
]


def _mk_rows() -> list[DetailRow]:
    """构造最小明细行：2 只股票 + 1 只 ETF（同账户）。"""
    return [
        DetailRow(
            account="证券",
            name="长江电力",
            code="600900",
            price=25.0,
            nav_date="2026-10-08",
            yesterday_close=24.5,
            shares=100.0,
            market_value=2500.0,
            cost=1000.0,
            profit=1500.0,
            profit_rate=1.5,
            today_profit=50.0,
        ),
        DetailRow(
            account="证券",
            name="贵州茅台",
            code="600519",
            price=1800.0,
            nav_date="2026-10-08",
            yesterday_close=1780.0,
            shares=50.0,
            market_value=90000.0,
            cost=10000.0,
            profit=80000.0,
            profit_rate=8.0,
            today_profit=1000.0,
        ),
        DetailRow(
            account="证券",
            name="沪深300ETF",
            code="510300",
            price=4.0,
            nav_date="2026-10-08",
            yesterday_close=3.9,
            shares=1000.0,
            market_value=4000.0,
            cost=3800.0,
            profit=200.0,
            profit_rate=0.05,
            today_profit=30.0,
        ),
    ]


def _mk_dicts(rows: list[DetailRow]) -> list[dict]:
    """由明细行构造 prepare 契约形态的明细字典（百分比收益率）。"""
    return [
        {
            "name": r.name,
            "code": r.code,
            "market_value": r.market_value,
            "cost": r.cost,
            "profit": r.profit,
            "profit_rate": (r.profit_rate or 0) * 100,
            "change_pct": 0.0,
            "nav_date": r.nav_date,
            "source_api": "mock",
            "shares": r.shares,
            "price": r.price,
            "channel": "场内",
        }
        for r in rows
    ]


def _mock_reporter() -> MagicMock:
    m = MagicMock()
    m.info = MagicMock()
    m.ok = MagicMock()
    m.warn = MagicMock()
    return m


class TestApplyReportAnonymization:
    """装配边界 pair 匿名入口。"""

    def test_off_mode_returns_same_objects(self):
        """off → 恒等返回（同一对象，零开销）。"""
        from src.python.report._report_helpers import apply_report_anonymization

        dicts, rows = _mk_dicts(_mk_rows()), _mk_rows()
        out_dicts, out_rows = apply_report_anonymization(dicts, rows, "off")
        assert out_dicts is dicts
        assert out_rows is rows

    def test_code_display_pair_shares_alias_and_keeps_codes(self):
        """code_display → 两容器名称同一代号，代码保留。"""
        from src.python.report._report_helpers import apply_report_anonymization

        rows = _mk_rows()
        dicts = _mk_dicts(rows)
        out_dicts, out_rows = apply_report_anonymization(dicts, rows, "code_display")

        for d, r in zip(out_dicts, out_rows, strict=True):
            assert d["name"] == r.name, "字典与行必须使用同一代号"
            assert d["code"] == r.code, "code_display 保留代码"
        assert out_dicts[0]["name"] == "品种A"
        assert out_dicts[1]["name"] == "品种B"
        # 数值不模糊
        assert out_dicts[0]["market_value"] == 2500.0
        assert out_rows[0].market_value == 2500.0

    def test_full_anonymous_blurs_values_and_keeps_code_keys(self):
        """full → 数值千位模糊、盈亏派生行内恒等、代码保留真值。"""
        from src.python.report._report_helpers import apply_report_anonymization

        rows = _mk_rows()
        dicts = _mk_dicts(rows)
        out_dicts, out_rows = apply_report_anonymization(dicts, rows, "full_anonymous")

        d0, r0 = out_dicts[0], out_rows[0]
        assert d0["code"] == "600900" and r0.code == "600900"
        assert d0["market_value"] == 2000.0  # round(2500/1000, 0)*1000（银行家舍入 → 2000）
        assert d0["cost"] == 1000.0
        assert d0["profit"] == 1000.0  # 派生：市值 − 成本（行内恒等）
        assert r0.market_value == 2000.0
        assert r0.profit == 1000.0
        assert r0.profit_rate == 1.0  # 小数契约：1000/1000
        assert d0["name"] == r0.name == "品种A"

    def test_summary_folds_dict_and_keeps_rows_identity(self):
        """summary → 字典折叠为大类聚合行，DetailRow 原样（渲染层拦截）。"""
        from src.python.report._report_helpers import apply_report_anonymization

        rows = _mk_rows()
        dicts = _mk_dicts(rows)
        out_dicts, out_rows = apply_report_anonymization(dicts, rows, "summary")

        assert out_rows is rows
        labels = [d["name"] for d in out_dicts]
        assert labels == ["股票汇总", "基金汇总"]
        # 合计守恒（折叠只是归并，不改总量）
        assert sum(d["market_value"] for d in out_dicts) == sum(d["market_value"] for d in dicts)


class TestFoldDetailRowsSummary:
    """明细渲染层折叠（账户组内按大类）。"""

    def test_folds_by_category_with_invariant_totals(self):
        from src.python.report._report_helpers import fold_detail_rows_summary

        rows = _mk_rows()
        folded = fold_detail_rows_summary(rows)

        assert all(isinstance(r, DetailRow) for r in folded)
        labels = [r.name for r in folded]
        assert labels == ["股票汇总", "基金汇总"]
        assert sum(r.market_value for r in folded) == sum(r.market_value for r in rows)
        assert sum(r.cost for r in folded) == sum(r.cost for r in rows)
        assert sum(r.profit for r in folded) == sum(r.profit for r in rows)
        for r in folded:
            assert r.code == ""

    def test_empty_rows_returns_empty(self):
        from src.python.report._report_helpers import fold_detail_rows_summary

        assert fold_detail_rows_summary([]) == []


class TestPrepareInjection:
    """prepare_report_data 装配边界注入（明细 pair + 约束块文本）。"""

    @patch("src.python.config.anonymizer.get_anonymization_mode", return_value="code_display")
    @patch("src.python.fetcher.index.fetch_indices")
    @patch("src.python.fetcher.index.fetch_us_indices")
    @patch("src.python.report.market_value._generate_details")
    def test_prepare_anonymizes_details_pair_and_constraint_block(
        self,
        mock_gen: MagicMock,
        mock_us: MagicMock,
        mock_a: MagicMock,
        mock_mode: MagicMock,
    ):
        mock_gen.return_value = _mk_rows()
        mock_a.return_value = {}
        mock_us.return_value = {}
        # 约束块构建期用真实名称渲染 → 注入后须按同一映射脱敏
        with patch(
            "src.python.report.orchestrator.compute_purchase_status_data",
            return_value={"constraint_block": "贵州茅台(600519) 限购 100 份/日", "available": True},
        ):
            from src.python.report.orchestrator import prepare_report_data

            prep = prepare_report_data(_SAMPLE_HOLDINGS, _mock_reporter(), {})

        # 明细字典与 DetailRow 同代号（pair 一致）
        dict_names = [d["name"] for d in prep["holdings_details"]]
        row_names = [r.name for r in prep["details"]]
        assert dict_names == row_names
        assert dict_names[0] == "品种A"
        # code_display 保留代码（键控链路）
        assert prep["holdings_details"][0]["code"] == "600900"
        # 约束块文本掩码（真名消失、代号出现）
        block = prep["purchase_status_data"]["constraint_block"]
        assert "贵州茅台" not in block
        assert "品种B" in block
        # 行动建议基于匿名明细同源构建
        assert prep["action_data"] is not None

    @patch("src.python.config.anonymizer.get_anonymization_mode", return_value="off")
    @patch("src.python.fetcher.index.fetch_indices")
    @patch("src.python.fetcher.index.fetch_us_indices")
    @patch("src.python.report.market_value._generate_details")
    def test_prepare_off_mode_keeps_real_names(
        self,
        mock_gen: MagicMock,
        mock_us: MagicMock,
        mock_a: MagicMock,
        mock_mode: MagicMock,
    ):
        mock_gen.return_value = _mk_rows()
        mock_a.return_value = {}
        mock_us.return_value = {}

        from src.python.report.orchestrator import prepare_report_data

        prep = prepare_report_data(_SAMPLE_HOLDINGS, _mock_reporter(), {})
        assert prep["holdings_details"][0]["name"] == "长江电力"
        assert prep["details"][0].name == "长江电力"


class TestExcelBasicGate:
    """Excel basic 路径：resolve_market_data 内部生成后就地匿名。"""

    @patch("src.python.config.anonymizer.get_anonymization_mode", return_value="code_display")
    def test_internal_generated_details_anonymized(self, mock_mode: MagicMock):
        from src.python.report.excel_market_data import resolve_market_data
        from src.python.report.progress import SilentProgressReporter

        gen = MagicMock(return_value=_mk_rows())
        modules = {"_generate_details": gen}
        data = resolve_market_data(_SAMPLE_HOLDINGS, None, modules, SilentProgressReporter())
        assert data["details"][0].name == "品种A"
        assert data["details"][0].code == "600900"

    @patch("src.python.config.anonymizer.get_anonymization_mode", return_value="code_display")
    def test_external_details_not_double_anonymized(self, mock_mode: MagicMock):
        """外部传入明细（prepare 已匿名）→ 不二次匿名（代号不再漂移）。"""
        from src.python.report.excel_market_data import resolve_market_data
        from src.python.report.progress import SilentProgressReporter

        pre_anon = [
            DetailRow(account="证券", name="品种A", code="600900", market_value=2500.0, cost=1000.0, profit=1500.0)
        ]
        gen = MagicMock(side_effect=AssertionError("外部明细不得触发内部生成"))
        data = resolve_market_data(_SAMPLE_HOLDINGS, pre_anon, {"_generate_details": gen}, SilentProgressReporter())
        assert data["details"][0].name == "品种A"


class TestHtmlGates:
    """HTML 路径：内部生成门 + 明细分组渲染拦截。"""

    @patch("src.python.config.anonymizer.get_anonymization_mode", return_value="code_display")
    @patch("src.python.report.html_renderers._generate_details")
    def test_market_value_section_internal_gen_anonymized(self, mock_gen: MagicMock, mock_mode: MagicMock):
        from src.python.report.html_renderers import _render_market_value_section
        from src.python.report.progress import SilentProgressReporter

        mock_gen.return_value = _mk_rows()
        details, totals = _render_market_value_section(_SAMPLE_HOLDINGS, None, "2026-10-08", SilentProgressReporter())
        assert details[0].name == "品种A"
        assert totals[0] == 96500.0  # 2500 + 90000 + 4000（code_display 不模糊）

    @patch("src.python.config.anonymizer.get_anonymization_mode", return_value="summary")
    def test_grouping_folds_summary_rows(self, mock_mode: MagicMock):
        from src.python.report.html_renderers import _render_account_grouping
        from src.python.report.progress import SilentProgressReporter

        rows = _mk_rows()
        accounts, totals = _render_account_grouping(rows, SilentProgressReporter())
        folded = accounts["证券"]
        assert [r.name for r in folded] == ["股票汇总", "基金汇总"]
        assert totals["证券"]["market_value"] == sum(r.market_value for r in rows)

    @patch("src.python.config.anonymizer.get_anonymization_mode", return_value="full_anonymous")
    def test_grouping_masks_code_display(self, mock_mode: MagicMock):
        """full → 明细行代码显示为 000XXX（字段层保留真值，渲染层兑现显示）。"""
        from src.python.report.html_renderers import _render_account_grouping
        from src.python.report.progress import SilentProgressReporter

        accounts, _ = _render_account_grouping(_mk_rows(), SilentProgressReporter())
        assert all(r.code == "000XXX" for r in accounts["证券"])

    @patch("src.python.config.anonymizer.get_anonymization_mode", return_value="code_display")
    def test_grouping_code_display_keeps_codes(self, mock_mode: MagicMock):
        from src.python.report.html_renderers import _render_account_grouping
        from src.python.report.progress import SilentProgressReporter

        accounts, _ = _render_account_grouping(_mk_rows(), SilentProgressReporter())
        assert [r.code for r in accounts["证券"]] == ["600900", "600519", "510300"]


class TestMaskWorkbookText:
    """Excel 产物清扫：字符串单元格替换、数值单元格不动。"""

    def test_masks_string_cells_and_skips_numeric(self):
        from openpyxl import Workbook

        from src.python.report.excel_writer import mask_workbook_text

        wb = Workbook()
        ws = wb.active
        ws["A1"] = "贵州茅台(600519)"
        ws["A2"] = 600519  # 数值单元格：即使等于代码也不动（无数字子串误伤）
        ws["A3"] = 42.5
        alias = {"贵州茅台": "品种B", "600519": "000XXX"}

        changed = mask_workbook_text(wb, alias)
        assert changed == 1
        assert ws["A1"].value == "品种B(000XXX)"
        assert ws["A2"].value == 600519
        assert ws["A3"].value == 42.5

    def test_empty_alias_map_is_noop(self):
        from openpyxl import Workbook

        from src.python.report.excel_writer import mask_workbook_text

        wb = Workbook()
        wb.active["A1"] = "贵州茅台"
        assert mask_workbook_text(wb, {}) == 0
        assert wb.active["A1"].value == "贵州茅台"


class TestNewsDisplayMask:
    """新闻关键词显示字段脱敏（取数链真值不变）。"""

    def test_masks_matched_and_enriched_labels(self):
        from src.python.report.news_correlation import _mask_news_display_fields

        items = [
            {
                "title": "贵州茅台三季度业绩",
                "matched_keywords": ["贵州茅台"],
                "enriched_keywords": [{"label": "贵州茅台", "type": "holding"}],
            }
        ]
        _mask_news_display_fields(items, _SAMPLE_HOLDINGS, "code_display")
        assert items[0]["matched_keywords"] == ["品种B"]
        assert items[0]["enriched_keywords"][0]["label"] == "品种B"
        # 检索真值（标题外部文本）不改写
        assert items[0]["title"] == "贵州茅台三季度业绩"

    def test_off_mode_noop(self):
        from src.python.report.news_correlation import _mask_news_display_fields

        items = [{"matched_keywords": ["贵州茅台"], "enriched_keywords": [{"label": "贵州茅台"}]}]
        _mask_news_display_fields(items, _SAMPLE_HOLDINGS, "off")
        assert items[0]["matched_keywords"] == ["贵州茅台"]
