"""持仓明细与分类 — 分类页签与整表写入（Excel 写入层）。

测试目标：
  - 分类页签写入与分类子列流量、分类申赎状态列
  - write_holdings_detail_sheet 整表入口与全零价格回归

运行：
  pytest src/test/unit/report/test_holdings_detail_categories.py -v
"""

from __future__ import annotations

import unittest
from unittest.mock import patch

from openpyxl import Workbook

from src.python.core.models import Holding
from src.python.report import holdings_detail_sheet as hds
from src.python.report.market_value import DetailRow
import pytest

pytestmark = [pytest.mark.unit, pytest.mark.unit_report, pytest.mark.usefixtures("offline_external_sources")]


# ═══════════════════════════════════════════════════════════
#  区块② 持仓分类汇总（由 test_category.py 迁入）
# ═══════════════════════════════════════════════════════════


class TestWriteCategorySheet(unittest.TestCase):
    """区块② 持仓分类汇总写入集成测试。"""

    def setUp(self):
        import openpyxl

        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active

    def test_single_holding(self):
        """单条持仓 → 分类表正确渲染。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(code="600900", account="证券", name="长江电力", market_value=2500, profit=1500, profit_rate=1.5)
        ]
        hds._write_category_block(self.ws, [h], details, None, 1)
        # 应该至少有一行标题 + 一行数据 + 合计行
        self.assertGreater(self.ws.max_row, 2)

    def test_empty_holdings(self):
        """空持仓 → 不崩溃。"""
        try:
            hds._write_category_block(self.ws, [], [], None, 1)
        except Exception as e:
            self.fail(f"空持仓分类区块写入不应抛异常: {e}")


class TestWriteCategoryFlowSubcolumns(unittest.TestCase):
    """持仓分类表成本流水子列（成本分档 / 分红累计）渲染测试。

    开关 `report_submodules.cost_lots` 对应 fund_flow_data 是否传入：
    None → 保持既有 10 列；非 None → 追加「成本分档」「分红累计」子列。
    """

    def setUp(self):
        import openpyxl

        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active

    def _call(self, holdings, details, fund_flow_data=None):
        # mock 分红 API 加载（避免真实网络请求，测试隔离）
        with patch("src.python.report.category._load_dividend_data", return_value=({}, True)):
            hds._write_category_block(self.ws, holdings, details, fund_flow_data, 1)

    def _headers(self):
        return [self.ws.cell(row=2, column=c).value for c in range(1, 13)]

    def _find_total_row(self):
        for r in range(1, self.ws.max_row + 1):
            if self.ws.cell(row=r, column=1).value == "总计":
                return r
        return None

    @staticmethod
    def _flow(low_shares=0.0, high_shares=0.0, div=0.0):
        return {
            "available": True,
            "cost_tiers": {
                "per_code": {
                    "600900": {
                        "low": {"shares": low_shares, "cost": low_shares * 9.0},
                        "high": {"shares": high_shares, "cost": high_shares * 11.0},
                        "unpriced": {"shares": 0.0, "cost": 0.0},
                    }
                }
            },
            "dividends": {"per_code": {"600900": div}},
        }

    def test_flow_subcolumns_header_when_enabled(self):
        """开关开启（fund_flow_data 非 None）时，表头追加「成本分档」「分红累计」列。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=1500,
                cost=1000,
                profit=500,
                profit_rate=0.5,
            )
        ]
        self._call([h], details, self._flow(low_shares=100, div=120.0))
        headers = self._headers()
        self.assertIn("成本分档", headers)
        self.assertIn("分红累计", headers)
        self.assertEqual(headers.index("成本分档") + 1, 11)
        self.assertEqual(headers.index("分红累计") + 1, 12)

    def test_flow_tier_low_and_div_value(self):
        """开关开启时数据行含「低成本」分档标签与分红累计数值。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=1500,
                cost=1000,
                profit=500,
                profit_rate=0.5,
            )
        ]
        self._call([h], details, self._flow(low_shares=100, div=120.0))
        # 数据行 = 表头下一行（row 3）：列 11 = 成本分档, 列 12 = 分红累计
        self.assertEqual(self.ws.cell(row=3, column=11).value, "低成本")
        self.assertEqual(self.ws.cell(row=3, column=12).value, 120.0)

    def test_flow_tier_high_label(self):
        """持仓批次成本价高于市价时渲染「高成本」档标签。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=900,
                cost=1000,
                profit=-100,
                profit_rate=-0.1,
            )
        ]
        self._call([h], details, self._flow(high_shares=100, div=0.0))
        self.assertEqual(self.ws.cell(row=3, column=11).value, "高成本")

    def test_flow_tier_mixed_label(self):
        """持仓批次横跨低/高两档时渲染「混合」档标签。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900", account="证券", name="长江电力", market_value=1000, cost=1000, profit=0, profit_rate=0.0
            )
        ]
        self._call([h], details, self._flow(low_shares=50, high_shares=50, div=0.0))
        self.assertEqual(self.ws.cell(row=3, column=11).value, "混合")

    def test_flow_total_row_dividend_sum(self):
        """开关开启时，分红累计在小计/总计行汇总（跨持仓累加）。"""
        flow = {
            "available": True,
            "cost_tiers": {
                "per_code": {
                    "600900": {
                        "low": {"shares": 100, "cost": 900},
                        "high": {"shares": 0, "cost": 0},
                        "unpriced": {"shares": 0, "cost": 0},
                    },
                    "600519": {
                        "low": {"shares": 10, "cost": 15000},
                        "high": {"shares": 0, "cost": 0},
                        "unpriced": {"shares": 0, "cost": 0},
                    },
                }
            },
            "dividends": {"per_code": {"600900": 120.0, "600519": 30.0}},
        }
        h1 = Holding("证券", "长江电力", "600900", 100, 10.0)
        h2 = Holding("证券", "贵州茅台", "600519", 10, 1500.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=1500,
                cost=1000,
                profit=500,
                profit_rate=0.5,
            ),
            DetailRow(
                code="600519",
                account="证券",
                name="贵州茅台",
                market_value=20000,
                cost=15000,
                profit=5000,
                profit_rate=0.33,
            ),
        ]
        self._call([h1, h2], details, flow)
        total_row = self._find_total_row()
        self.assertIsNotNone(total_row)
        self.assertEqual(self.ws.cell(row=total_row, column=12).value, 150.0)

    def test_no_flow_subcolumns_when_disabled(self):
        """开关关闭（fund_flow_data=None）时保持既有 10 列输出，无成本流水子列。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=1500,
                cost=1000,
                profit=500,
                profit_rate=0.5,
            )
        ]
        self._call([h], details, None)
        headers = self._headers()
        self.assertNotIn("成本分档", headers)
        self.assertNotIn("分红累计", headers)
        # 既有表头 10 列保持不变（前 10 列与改造前一致）
        self.assertEqual(headers[:10][-1], "年均股息率")


@pytest.mark.data

# ═══════════════════════════════════════════════════════════
#  write_holdings_detail_sheet — 合并章节写入器（区块① + 区块②）
# ═══════════════════════════════════════════════════════════


class TestWriteHoldingsDetailSheet(unittest.TestCase):
    """合并章节写入器：单页签承载两个区块（持仓明细与分类）。"""

    def setUp(self):
        import openpyxl

        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.details = [
            hds.DetailRow(
                account="A",
                name="贵州茅台",
                code="600519",
                price=1500.0,
                nav_date="",
                yesterday_close=1480.0,
                price_type="场内收盘价(T)",
                premium="--",
                shares=100.0,
                market_value=150000.0,
                cost=140000.0,
                profit=10000.0,
                profit_rate=0.0714,
                today_profit=2000.0,
                source="腾讯财经",
            ),
            hds.DetailRow(
                account="A",
                name="某基金",
                code="040046",
                price=1.234,
                nav_date="2026-09-14",
                yesterday_close=1.22,
                price_type="官方净值(T-1)",
                premium="--",
                shares=1000.0,
                market_value=1234.0,
                cost=1200.0,
                profit=34.0,
                profit_rate=0.028,
                today_profit=14.0,
                source="基金净值",
            ),
        ]
        self.holdings = [
            Holding(code="600519", name="贵州茅台", shares=100.0, cost_price=1400.0, account="A"),
            Holding(code="040046", name="某基金", shares=1000.0, cost_price=1.2, account="A"),
        ]

    def test_title_is_merged_sheet_name(self):
        """首行为合并章节显示名「持仓明细与分类」。"""
        from src.python.core.registry import get_report_sheet_name

        hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details)
        self.assertEqual(self.ws["A1"].value, get_report_sheet_name("holdings_detail"))
        self.assertEqual(self.ws["A1"].value, "持仓明细与分类")

    def test_two_blocks_present_in_one_sheet(self):
        """同一工作表内含两个区块小节标题（市值明细 / 分类汇总）。"""
        hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details)
        col_a = [self.ws.cell(row=r, column=1).value for r in range(1, self.ws.max_row + 1)]
        self.assertIn("一、市值核算明细", col_a)
        self.assertIn("二、持仓分类汇总", col_a)
        # 区块①表头与区块②表头各自出现（同页签两表）
        all_vals = [self.ws.cell(row=r, column=c).value for r in range(1, self.ws.max_row + 1) for c in range(1, 11)]
        self.assertIn("账户", all_vals)
        self.assertIn("资产属性", all_vals)

    def test_returns_market_summary_tuple(self):
        """返回值与原市值明细写入器一致：(总市值, 总成本, 总盈亏, 本日总盈亏, 明细行)。"""
        result = hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details)
        self.assertEqual(len(result), 5)
        self.assertAlmostEqual(result[0], 151234.0)
        self.assertAlmostEqual(result[1], 141200.0)
        self.assertAlmostEqual(result[2], 10034.0)
        self.assertAlmostEqual(result[3], 2014.0)
        self.assertEqual(len(result[4]), 2)

    def test_block_values_equal_standalone_blocks(self):
        """内容等价：合并页签的两个区块行值 == 各自独立写入（start_row=1）的结果。"""
        import openpyxl

        hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details)

        wb2 = openpyxl.Workbook()
        ws_mv = wb2.active
        hds._write_market_value_block(ws_mv, self.details, None, 1)
        ws_cat = wb2.create_sheet()
        hds._write_category_block(ws_cat, self.holdings, self.details, None, 1)

        # 区块①：跳过合并页签新增的区块小节标题行（第 2 行），其余逐行相等
        mv_merged = [
            [self.ws.cell(row=r + 1, column=c).value for c in range(1, 17)] for r in range(1, ws_mv.max_row + 1)
        ]
        mv_standalone = [[ws_mv.cell(row=r, column=c).value for c in range(1, 17)] for r in range(1, ws_mv.max_row + 1)]
        self.assertEqual(mv_merged[2:], mv_standalone[2:])  # 表头及数据行逐格相等

        # 区块②：分类行数据（资产属性/投资分类/名称/代码/市值）与独立写入一致
        cat_merged = [
            [self.ws.cell(row=r, column=c).value for c in range(1, 11)]
            for r in range(1, self.ws.max_row + 1)
            if self.ws.cell(row=r, column=1).value in ("股票", "基金", "现金", "债券", "其他")
        ]
        cat_standalone = [
            [ws_cat.cell(row=r, column=c).value for c in range(1, 11)]
            for r in range(1, ws_cat.max_row + 1)
            if ws_cat.cell(row=r, column=1).value in ("股票", "基金", "现金", "债券", "其他")
        ]
        self.assertEqual(cat_merged, cat_standalone)
        self.assertTrue(cat_merged, "分类区块应至少有一行分类明细")

    def test_flow_columns_apply_to_both_blocks(self):
        """成本流水开关开启时，两个区块各自追加子列（区块①「资金加权成本」/ 区块②「成本分档」）。"""
        flow = {"available": True, "approximate": False, "cost_tiers": {"per_code": {}}, "dividends": {"per_code": {}}}
        hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details, fund_flow_data=flow)
        all_vals = [self.ws.cell(row=r, column=c).value for r in range(1, self.ws.max_row + 1) for c in range(1, 17)]
        self.assertIn("资金加权成本", all_vals)
        self.assertIn("成本分档", all_vals)
        self.assertIn("分红累计", all_vals)


class TestAllZeroPriceRegression(unittest.TestCase):
    """缺陷回归：行情全零时提示行**整行合并**，数据必须自其后一行写入。

    现场（logs/app.log 2026-09-16 场景测试暴露）：提示行 `merge_cells` 后仍以该行为
    数据起点 → `AttributeError: 'MergedCell' object attribute 'value' is read-only`，
    导致整份 Excel 报告生成失败。既有实现（合并前的 market_value_sheet）同样有此缺陷，
    批次② 合并时原样带入。
    """

    def _ws(self):
        import openpyxl

        wb = openpyxl.Workbook()
        return wb.active

    def test_all_zero_price_writes_warning_and_data_without_crash(self):
        ws = self._ws()
        zero_details = [
            hds.DetailRow(
                account="A", name="贵州茅台", code="600519", price=0.0, shares=100.0, market_value=0.0, cost=140000.0
            ),
            hds.DetailRow(
                account="A", name="某基金", code="040046", price=0.0, shares=1000.0, market_value=0.0, cost=1200.0
            ),
        ]
        holdings = [
            Holding(code="600519", name="贵州茅台", shares=100.0, cost_price=1400.0, account="A"),
            Holding(code="040046", name="某基金", shares=1000.0, cost_price=1.2, account="A"),
        ]
        # 不得抛异常（缺陷现场在此崩溃）
        hds.write_holdings_detail_sheet(ws, holdings, zero_details)

        flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertTrue(any("行情数据全部不可用" in v for v in flat), "应写入全零提示行")
        self.assertIn("贵州茅台", flat, "数据行应写在提示行之后")
        self.assertIn("A 小计", flat)

    def test_all_zero_price_indicator_block_order(self):
        """区块① 中提示行之后的数据行号必须大于提示行行号（防写回合并区）。"""
        ws = self._ws()
        zero_details = [
            hds.DetailRow(
                account="A", name="贵州茅台", code="600519", price=0.0, shares=100.0, market_value=0.0, cost=140000.0
            )
        ]
        hds.write_holdings_detail_sheet(ws, [], zero_details)
        warn_row = next(
            r for r in range(1, ws.max_row + 1) if "行情数据全部不可用" in str(ws.cell(row=r, column=1).value or "")
        )
        data_row = next((r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "贵州茅台"), None)
        self.assertIsNotNone(data_row, "全零场景仍应写出数据行")
        self.assertGreater(data_row, warn_row, "数据行必须在提示行之后")


# ═══════════════════════════════════════════════════════════
#  区块②申购状态条件列（与区块①同开关同判据，fund_purchase_limit）
# ═══════════════════════════════════════════════════════════


class TestCategoryPurchaseColumn(unittest.TestCase):
    """区块②「申购状态」条件列（purchase_status_data 契约 → 分类汇总末列）。

    判据单源：列可见性 = purchase_column_visible、文案 = format_purchase_status_cell、
    时效 = stale_level —— 与区块①同一套原语，开关联动 `fund_purchase_limit` 不新增。
    """

    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.holding = Holding(account="证券账户", name="电池ETF", code="561910", shares=100, cost_price=1.0)
        self.detail = hds.DetailRow(
            account="证券账户",
            name="电池ETF",
            code="561910",
            price=10.0,
            nav_date="2026-06-26",
            yesterday_close=9.5,
            price_type="场内收盘价(T)",
            premium="--",
            shares=100.0,
            market_value=1000.0,
            cost=100.0,
            profit=900.0,
            profit_rate=9.0,
            today_profit=50.0,
            source="腾讯财经",
            source_api="tencent",
        )

    @staticmethod
    def _contract(available: bool = True) -> dict:
        return {
            "available": available,
            "reason": None if available else "全链路失败",
            "rows": {
                "561910": {
                    "purchase_status": "限大额",
                    "redemption_status": "开放赎回",
                    "next_open_date": "",
                    "daily_limit": 100.0,
                    "min_purchase": 10.0,
                }
            },
            "fetched_at": "2099-01-01T08:00:00+08:00",
            "source": "天天基金",
        }

    def _write(self, contract: dict | None = None, flow: dict | None = None) -> None:
        # mock 分红 API 加载（测试隔离，避免真实网络请求）
        with patch("src.python.report.category._load_dividend_data", return_value=({}, True)):
            hds._write_category_block(self.ws, [self.holding], [self.detail], flow, 1, contract)

    @staticmethod
    def _headers(ws, count: int, row: int = 2) -> list:
        return [ws.cell(row=row, column=c).value for c in range(1, count + 1)]

    def test_header_appended_when_available(self):
        """available=True → 区块②表头第 11 列「申购状态」，原 10 列保持不变。"""
        self._write(self._contract())
        headers = self._headers(self.ws, 11)
        self.assertEqual(headers[9], "年均股息率")
        self.assertEqual(headers[10], "申购状态")

    def test_cell_text_from_single_source(self):
        """分类明细行单元格文案与单源函数逐字一致（level 取值同式对齐）。"""
        from src.python.report.purchase_status import format_purchase_status_cell, stale_level

        contract = self._contract()
        self._write(contract)
        expected = format_purchase_status_cell(contract, "561910", stale_level(contract["fetched_at"]))
        self.assertEqual(self.ws.cell(row=3, column=11).value, expected)
        self.assertIn("限大额", expected)

    def test_subtotal_and_total_cell_blank(self):
        """分组小计行与总计行申购状态列留空（非可聚合指标，与区块①同语义）。"""
        self._write(self._contract())
        for row in (4, 5):  # 行3 明细 / 行4 分组小计 / 行5 总计
            self.assertIn(self.ws.cell(row=row, column=11).value, (None, ""))

    def test_hidden_when_contract_none(self):
        """契约缺席（开关关）→ 与既有基线一致的 10 列，无申购状态。"""
        self._write(None)
        headers = self._headers(self.ws, 11)
        self.assertEqual(headers[9], "年均股息率")
        self.assertEqual(headers[10], None)
        self.assertNotIn("申购状态", [h for h in headers if h])

    def test_hidden_when_unavailable(self):
        """available=False（全链失败）→ 静默隐列：保持既有 10 列。"""
        self._write(self._contract(available=False))
        headers = self._headers(self.ws, 11)
        self.assertNotIn("申购状态", [h for h in headers if h])

    def test_flow_and_purchase_coexist(self):
        """成本流水 + 申购状态同时开启 → 13 列，申购状态在末列。"""
        flow = {"available": True, "cost_tiers": {"per_code": {}}, "dividends": {"per_code": {}}}
        self._write(self._contract(), flow=flow)
        headers = self._headers(self.ws, 13)
        self.assertEqual(headers[10], "成本分档")
        self.assertEqual(headers[11], "分红累计")
        self.assertEqual(headers[12], "申购状态")
        self.assertIn("限大额", self.ws.cell(row=3, column=13).value)
