"""LLM 模块输出质量分级单元测试。

覆盖 A~F 分级口径、横幅注入位置与开关行为，并锁定「必需章节标记」
与 LLM System Prompt 的一致性（防止提示词改章节后分级静默误判）。
"""

from __future__ import annotations

import pytest

pytestmark = [pytest.mark.unit, pytest.mark.unit_report]

from unittest.mock import MagicMock

from src.python.report.llm_content import _FACT_CHECK_FAIL_RE, _FACT_CHECK_PASS_RE
from src.python.report.llm_quality import (
    GRADE_A,
    GRADE_B,
    GRADE_C,
    GRADE_D,
    GRADE_F,
    TRIGGER_EMPTY,
    TRIGGER_MISSING_SECTIONS,
    TRIGGER_OK,
    TRIGGER_PLACEHOLDER,
    TRIGGER_THIN,
    TRIGGER_TOO_SHORT,
    _LENGTH_THRESHOLDS,
    _REQUIRED_MARKERS,
    apply_quality_banners,
    build_quality_banner,
    grade_module,
    grade_modules,
)

# 健康体检报告所需章节齐备的正文（篇幅远超参考篇幅）
_HEALTH_FULL = (
    "综合评分 总分：84/100 评级：良\n"
    "一、风险分散度（82/100）\n"
    "二、流动性（80/100）\n"
    "三、收益合理性（85/100）\n"
    "四、成本结构（88/100）\n"
    "五、数据质量（90/100）\n"
    "改进建议\n" + "评分依据与说明。" * 200
)

_PENETRATION_FULL = (
    "行业集中度分析 前5大行业及占比\n"
    "品种集中度分析 TOP10 底层资产\n"
    "国别/币种暴露 A股/港股/美股\n"
    "综合建议 2-3条调整建议\n" + "数据分析说明。" * 200
)


def _feature_on(monkeypatch) -> None:
    """在当前用例内开启模块级质量分级开关。"""
    from src.python.config import features as feat

    monkeypatch.setitem(feat.FEATURE_FLAGS, "module_quality_gate", True)


def _feature_off(monkeypatch) -> None:
    from src.python.config import features as feat

    monkeypatch.setitem(feat.FEATURE_FLAGS, "module_quality_gate", False)


@pytest.mark.unit
class TestGradeModule:
    """grade_module 分级口径测试。"""

    def test_healthy_health_check_grades_a(self):
        """章节齐备且篇幅达标 → A。"""
        q = grade_module("health_check", _HEALTH_FULL)
        assert q.grade == GRADE_A
        assert q.trigger == TRIGGER_OK

    def test_healthy_penetration_grades_a(self):
        """穿透深度分析章节齐备 → A。"""
        assert grade_module("penetration_deep", _PENETRATION_FULL).grade == GRADE_A

    def test_none_grades_f(self):
        """内容为 None → F（正文为空）。"""
        q = grade_module("health_check", None)
        assert q.grade == GRADE_F
        assert q.trigger == TRIGGER_EMPTY

    def test_blank_html_grades_f(self):
        """只有标签没有正文 → F。"""
        q = grade_module("global_macro", "<p>   </p><br/>")
        assert q.grade == GRADE_F
        assert q.trigger == TRIGGER_EMPTY

    def test_placeholder_grades_d(self):
        """命中降级占位模板 → D。"""
        from src.python.llm.fallback import get_placeholder_text

        q = grade_module("penetration_deep", get_placeholder_text("penetration_deep"))
        assert q.grade == GRADE_D
        assert q.trigger == TRIGGER_PLACEHOLDER

    def test_truncation_marker_grades_d(self):
        """命中截断标记 → D，即使篇幅很长。"""
        from src.python.llm.api_base import TRUNCATION_MARKER

        q = grade_module("health_check", _HEALTH_FULL + TRUNCATION_MARKER + "！上下文上限不足")
        assert q.grade == GRADE_D
        assert q.trigger == TRIGGER_PLACEHOLDER

    def test_too_short_grades_d(self):
        """正文短于模块下限 → D。"""
        min_chars = _LENGTH_THRESHOLDS["health_check"][0]
        q = grade_module("health_check", "<p>" + "体检" * (min_chars // 4) + "</p>")
        assert q.grade == GRADE_D
        assert q.trigger == TRIGGER_TOO_SHORT

    def test_missing_sections_grades_c(self):
        """篇幅达标但缺必需章节 → C，且原因列出缺失项。"""
        body = "<p>综合评分 总分：84/100 " + "评分依据。" * 200 + "</p>"
        q = grade_module("health_check", body)
        assert q.grade == GRADE_C
        assert q.trigger == TRIGGER_MISSING_SECTIONS
        assert "风险分散度" in q.reason
        assert "综合评分" not in q.reason  # 已有章节不应被列为缺失

    def test_thin_but_complete_grades_b(self):
        """章节齐备、篇幅介于下限与参考篇幅之间 → B。"""
        min_chars, comfort_chars = _LENGTH_THRESHOLDS["health_check"]
        markers = "".join(_REQUIRED_MARKERS["health_check"])
        # 补齐到两阈值中点，确保严格落在 (下限, 参考篇幅) 区间内
        pad = (min_chars + comfort_chars) // 2 - len(markers)
        q = grade_module("health_check", "<p>" + markers + "填" * pad + "</p>")
        assert q.grade == GRADE_B
        assert q.trigger == TRIGGER_THIN

    def test_prose_module_not_penalised_for_missing_headings(self):
        """纯散文模块（global_macro）无必需章节，长文即 A。"""
        assert grade_module("global_macro", "<p>" + "宏观分析。" * 300 + "</p>").grade == GRADE_A

    def test_expert_review_not_penalised_for_missing_headings(self):
        """智囊团复盘章节随辩论模式变化，不做章节判定（只判篇幅）。"""
        assert grade_module("expert_review", "<p>" + "复盘分析。" * 400 + "</p>").grade == GRADE_A


@pytest.mark.unit
class TestRequiredMarkersMatchPrompts:
    """必需章节标记与 System Prompt 一致性测试。

    标记是启发式判定，一旦提示词的章节小标题与标记不同步，分级会静默
    误判（把健康输出评成 C）。此处直接比对提示词常量，二者不同步即失败。
    """

    def test_health_check_markers_present_in_system_prompt(self):
        from src.python.llm.prompts_core import _SYSTEM_HEALTH_CHECK

        for marker in _REQUIRED_MARKERS["health_check"]:
            assert marker in _SYSTEM_HEALTH_CHECK, f"标记「{marker}」已不在体检提示词的章节规定中"

    def test_penetration_markers_present_in_system_prompt(self):
        from src.python.llm.prompts_core import _SYSTEM_PENETRATION_DEEP

        for marker in _REQUIRED_MARKERS["penetration_deep"]:
            assert marker in _SYSTEM_PENETRATION_DEEP, f"标记「{marker}」已不在穿透提示词的章节规定中"

    def test_marker_modules_have_length_thresholds(self):
        """配置了必需章节的模块必须同时配置长度阈值，避免阈值缺失走默认值。"""
        for module_key in _REQUIRED_MARKERS:
            assert module_key in _LENGTH_THRESHOLDS


@pytest.mark.unit
class TestGradeModules:
    """grade_modules 批量分级测试。"""

    def test_skips_none_modules(self):
        """未生成的模块（None）不产生评级条目。"""
        results = grade_modules((_HEALTH_FULL, None, _HEALTH_FULL, None))
        assert [q.module_key for q in results] == ["global_macro", "health_check"]

    def test_short_tuple_handled(self):
        """元组短于模块数时不越界。"""
        assert grade_modules((_HEALTH_FULL,))[0].module_key == "global_macro"

    def test_empty_input(self):
        assert grade_modules(()) == []


@pytest.mark.unit
class TestApplyQualityBanners:
    """apply_quality_banners 横幅注入与开关行为测试。"""

    def test_flag_off_returns_same_object(self, monkeypatch):
        """开关关闭时原样返回同一对象（零开销、零改写）。"""
        _feature_off(monkeypatch)
        content = (None, None, "<p>" + "体检" * 900 + "</p>", None)
        assert apply_quality_banners(content) is content

    def test_flag_on_returns_new_tuple(self, monkeypatch):
        """开关开启且内容无问题时返回新元组，内容等值。"""
        _feature_on(monkeypatch)
        content = (None, None, _HEALTH_FULL, None)
        out = apply_quality_banners(content)
        assert out == content
        assert out is not content

    def test_banner_prepended_for_advisory_grade(self, monkeypatch):
        """低评级（内容在但缺章节）→ 横幅置于内容头部。"""
        _feature_on(monkeypatch)
        body = "<p>" + "评分说明。" * 300 + "</p>"  # 缺章节、篇幅足够 → C
        out = apply_quality_banners((None, None, body, None))
        assert out[2].startswith("<p style=")
        assert "【内容质量提示】" in out[2]
        assert out[2].endswith("</p>")
        assert body in out[2]

    def test_placeholder_not_bannered(self, monkeypatch):
        """降级占位已有醒目提示，不再叠加质量横幅。"""
        _feature_on(monkeypatch)
        from src.python.llm.fallback import get_placeholder_text

        placeholder = get_placeholder_text("health_check")
        assert apply_quality_banners((None, None, placeholder, None))[2] == placeholder

    def test_healthy_content_untouched(self, monkeypatch):
        """A 级内容不加任何横幅（健康输出零噪音）。"""
        _feature_on(monkeypatch)
        out = apply_quality_banners((_HEALTH_FULL, None, _HEALTH_FULL, None))
        assert "内容质量提示" not in out[0]
        assert "内容质量提示" not in out[2]

    def test_none_modules_stay_none(self, monkeypatch):
        """None 模块不被替换为占位（占位由降级链路负责）。"""
        _feature_on(monkeypatch)
        assert apply_quality_banners((None, None, None, None)) == (None, None, None, None)

    def test_non_contiguous_none_positions(self, monkeypatch):
        """只有末位模块有内容时，横幅落在末位。"""
        _feature_on(monkeypatch)
        body = "<p>" + "数据说明。" * 300 + "</p>"
        out = apply_quality_banners((None, None, None, body))
        assert out[:3] == (None, None, None)
        assert "【内容质量提示】" in out[3]

    def test_reporter_warned_once_for_advisory(self, monkeypatch):
        """存在低评级模块时向 reporter 汇总告警一次。"""
        _feature_on(monkeypatch)
        reporter = MagicMock()
        body = "<p>" + "评分说明。" * 300 + "</p>"
        apply_quality_banners((None, None, body, None), reporter)
        reporter.warn.assert_called_once()
        assert "持仓体检报告" in reporter.warn.call_args.args[0]

    def test_no_reporter_warn_for_healthy(self, monkeypatch):
        """全部健康时不打扰用户。"""
        _feature_on(monkeypatch)
        reporter = MagicMock()
        apply_quality_banners((_HEALTH_FULL, None, _HEALTH_FULL, None), reporter)
        reporter.warn.assert_not_called()

    def test_reporter_none_does_not_crash(self, monkeypatch):
        """未传 reporter 时仍可完成注入。"""
        _feature_on(monkeypatch)
        body = "<p>" + "评分说明。" * 300 + "</p>"
        assert "【内容质量提示】" in apply_quality_banners((None, None, body, None), None)[2]


@pytest.mark.unit
class TestQualityBannerFormat:
    """横幅 HTML 与既有内容块的兼容性测试。"""

    def test_banner_not_mistaken_for_fact_check_block(self, monkeypatch):
        """横幅不得命中事实校验摘要正则（否则 Excel 会按校验行着色）。"""
        banner = build_quality_banner(grade_module("health_check", "<p>" + "x" * 800 + "</p>"))
        text = "<p>" + banner + "</p>"
        assert not _FACT_CHECK_PASS_RE.search("【内容质量提示】")
        assert not _FACT_CHECK_FAIL_RE.search("【内容质量提示】")
        assert text  # 横幅本身可渲染

    def test_banner_carries_grade_and_reason(self):
        """横幅展示评级与原因，读者无需查表。"""
        q = grade_module("health_check", "<p>" + "评分说明。" * 300 + "</p>")
        banner = build_quality_banner(q)
        assert q.grade in banner
        assert "缺失必需章节" in banner


@pytest.mark.unit
class TestBannerExcelCarrier:
    """横幅经 Excel 页签渲染的载体契约测试。

    横幅复用模块 HTML 字符串作载体，故须确认它落在 Excel 侧被当作普通
    段落处理（普通字体、换行开启），而不是被误判为事实校验告警行。
    """

    def _first_content_row(self, content: str) -> tuple[str, object]:
        from openpyxl import Workbook

        from src.python.report.llm_content import _write_content_sheet

        wb = Workbook()
        ws = wb.create_sheet()
        _write_content_sheet(ws, "持仓体检报告", content)
        cell = ws.cell(row=2, column=1)
        return str(cell.value), cell

    def test_banner_is_first_content_row_with_normal_font(self, monkeypatch):
        from src.python.report.styles import CONTENT_FONT

        _feature_on(monkeypatch)
        body = "<p>" + "评分说明。" * 300 + "</p>"
        content = apply_quality_banners((None, None, body, None))[2]

        value, cell = self._first_content_row(content)
        assert value.startswith("【内容质量提示】")
        assert cell.font.color.rgb == CONTENT_FONT.color.rgb
        assert cell.alignment.wrap_text is True

    def test_body_still_present_after_banner(self, monkeypatch):
        """横幅之后的正文段落完整保留（不吞内容）。"""
        _feature_on(monkeypatch)
        body = "<p>" + "评分说明。" * 300 + "</p>"
        content = apply_quality_banners((None, None, body, None))[2]

        from openpyxl import Workbook

        from src.python.report.llm_content import _write_content_sheet

        wb = Workbook()
        ws = wb.create_sheet()
        _write_content_sheet(ws, "持仓体检报告", content)
        rows = [ws.cell(row=r, column=1).value for r in range(2, 6)]
        assert any(isinstance(v, str) and v.startswith("评分说明") for v in rows)
