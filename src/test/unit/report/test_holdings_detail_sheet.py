"""持仓明细与分类 — Excel 写入层单元测试（区块①市值明细）。

测试目标：
  - _detail_to_row_values   — 行值转换
  - _mv_num_formats         — 区块①格式列表
  - _apply_profit_colors    — 盈亏着色
  - _apply_price_type_colors — 取价方式着色
  - _write_account_groupings — 账户分组写入
  - _write_market_value_block — 区块①写入（mock 内部函数）

通过直接导入 holdings_detail_sheet 模块引用写入层函数。
"""

from __future__ import annotations

import unittest
from typing import Any
from unittest.mock import MagicMock, patch

from openpyxl import Workbook

from src.python.core.models import Holding
from src.python.report import holdings_detail_sheet as hds
from src.python.report.market_value import DetailRow
import pytest

pytestmark = [pytest.mark.unit, pytest.mark.unit_report]


# ═══════════════════════════════════════════════════════════
#  _detail_to_row_values
# ═══════════════════════════════════════════════════════════


class TestDetailToRowValues(unittest.TestCase):
    """测试 _detail_to_row_values 行值转换。"""

    def test_full_detail_row(self):
        """完整 DetailRow → 15 个字段的列表。"""
        d = hds.DetailRow(
            account="证券账户",
            name="电池ETF",
            code="561910",
            price=10.5,
            nav_date="2026-06-26",
            yesterday_close=10.0,
            price_type="场内收盘价(T)",
            premium="--",
            shares=1000.0,
            market_value=10500.0,
            cost=1000.0,
            profit=9500.0,
            profit_rate=9.5,
            today_profit=500.0,
            source="腾讯财经",
            source_api="tencent",
        )
        vals = hds._detail_to_row_values(d)
        self.assertEqual(len(vals), 15)
        self.assertEqual(vals[0], "证券账户")
        self.assertEqual(vals[1], "电池ETF")
        self.assertEqual(vals[2], "561910")
        self.assertEqual(vals[3], 10.5)
        self.assertEqual(vals[4], "2026-06-26")
        self.assertEqual(vals[5], 10.0)
        self.assertEqual(vals[6], "场内收盘价(T)")
        self.assertEqual(vals[7], "--")
        self.assertEqual(vals[8], 1000.0)
        self.assertEqual(vals[9], 10500.0)
        self.assertEqual(vals[10], 1000.0)
        self.assertEqual(vals[11], 9500.0)
        self.assertEqual(vals[12], 9.5)
        self.assertEqual(vals[13], 500.0)
        self.assertEqual(vals[14], "腾讯财经")

    def test_default_row(self):
        """默认 DetailRow → 空字符串/0 值列表。"""
        d = hds.DetailRow()
        vals = hds._detail_to_row_values(d)
        self.assertEqual(len(vals), 15)
        self.assertEqual(vals[0], "")
        self.assertEqual(vals[3], 0.0)
        self.assertEqual(vals[7], "--")
        self.assertIsNone(vals[12])


# ═══════════════════════════════════════════════════════════
#  _num_formats
# ═══════════════════════════════════════════════════════════


class TestNumFormats(unittest.TestCase):
    """测试 _num_formats 返回正确长度的格式列表。"""

    def test_length(self):
        """返回 15 个格式。"""
        fmts = hds._mv_num_formats()
        self.assertEqual(len(fmts), 15)

    def test_price_format(self):
        """第 4 个为价格格式。"""
        fmts = hds._mv_num_formats()
        self.assertEqual(fmts[3], "#,##0.0000")

    def test_money_format(self):
        """市值/成本/盈亏列金额格式。"""
        fmts = hds._mv_num_formats()
        self.assertEqual(fmts[9], "#,##0.00")
        self.assertEqual(fmts[10], "#,##0.00")
        self.assertEqual(fmts[11], "#,##0.00")

    def test_percent_format(self):
        """收益率列百分比格式。"""
        fmts = hds._mv_num_formats()
        self.assertEqual(fmts[12], "0.00%")

    def test_shares_format(self):
        """份额列格式。"""
        fmts = hds._mv_num_formats()
        self.assertEqual(fmts[8], "#,##0.00")


# ═══════════════════════════════════════════════════════════
#  _apply_profit_colors
# ═══════════════════════════════════════════════════════════


class TestApplyProfitColors(unittest.TestCase):
    """测试 _apply_profit_colors 盈亏列着色（mock ws / profit_font）。"""

    def _make_cell(self, value: Any = None) -> MagicMock:
        cell = MagicMock()
        cell.value = value
        return cell

    @patch("src.python.report.holdings_detail_sheet.profit_font")
    def test_positive_profit(self, mock_pf):
        """正盈亏 → 调用 profit_font(正数)。"""
        mock_pf.side_effect = lambda v: f"font_for_{v}"
        ws = MagicMock()
        ws.cell.side_effect = lambda row, column: self._make_cell(500.0)
        hds._apply_profit_colors(ws, 3, 3, profit_col=12, rate_col=13, today_col=14)
        mock_pf.assert_any_call(500.0)
        self.assertEqual(mock_pf.call_count, 3)

    @patch("src.python.report.holdings_detail_sheet.profit_font")
    def test_negative_profit(self, mock_pf):
        """负盈亏 → 调用 profit_font(负数)。"""
        mock_pf.side_effect = lambda v: f"font_for_{v}"
        ws = MagicMock()
        ws.cell.side_effect = lambda row, column: self._make_cell(-300.0)
        hds._apply_profit_colors(ws, 3, 3, profit_col=12, rate_col=13, today_col=14)
        mock_pf.assert_any_call(-300.0)

    @patch("src.python.report.holdings_detail_sheet.profit_font")
    def test_zero_profit(self, mock_pf):
        """零盈亏 → 调用 profit_font(0)。"""
        mock_pf.side_effect = lambda v: f"font_for_{v}"
        ws = MagicMock()
        ws.cell.side_effect = lambda row, column: self._make_cell(0.0)
        hds._apply_profit_colors(ws, 3, 3, profit_col=12, rate_col=13, today_col=14)
        mock_pf.assert_any_call(0.0)

    @patch("src.python.report.holdings_detail_sheet.profit_font")
    def test_non_numeric_skipped(self, mock_pf):
        """非数字值（字符串）→ 不设置字体。"""
        ws = MagicMock()
        values = {"12": "亏损", "14": "盈利"}

        def cell_side_effect(row, column):
            cell = MagicMock()
            cell.value = values.get(str(column))
            return cell

        ws.cell.side_effect = cell_side_effect
        hds._apply_profit_colors(ws, 3, 4, profit_col=12, rate_col=13, today_col=14)
        mock_pf.assert_not_called()

    @patch("src.python.report.holdings_detail_sheet.profit_font")
    def test_none_value_skipped(self, mock_pf):
        """None 值 → 不设置字体。"""
        ws = MagicMock()
        ws.cell.return_value = self._make_cell(None)
        hds._apply_profit_colors(ws, 3, 4, profit_col=12, rate_col=13, today_col=14)
        mock_pf.assert_not_called()

    @patch("src.python.report.holdings_detail_sheet.profit_font")
    def test_rate_col_float(self, mock_pf):
        """收益率列为 float → 调用 profit_font。"""
        mock_pf.side_effect = lambda v: f"font_for_{v}"
        ws = MagicMock()
        ws.cell.side_effect = lambda row, column: self._make_cell(0.05)
        hds._apply_profit_colors(ws, 3, 3, profit_col=12, rate_col=13, today_col=14)
        mock_pf.assert_called_with(0.05)

    @patch("src.python.report.holdings_detail_sheet.profit_font")
    def test_rate_col_not_float(self, mock_pf):
        """收益率列非 float（如字符串含 %）→ 不设置字体。"""
        ws = MagicMock()
        ws.cell.side_effect = lambda row, column: self._make_cell("5.00%")
        hds._apply_profit_colors(ws, 3, 4, profit_col=12, rate_col=13, today_col=14)
        mock_pf.assert_not_called()

    @patch("src.python.report.holdings_detail_sheet.profit_font")
    def test_multiple_rows(self, mock_pf):
        """多行数据 → 每行都着色。"""
        mock_pf.side_effect = lambda v: f"font_for_{v}"
        ws = MagicMock()
        row_values = {3: 100.0, 4: -50.0, 5: 200.0}

        def cell_side_effect(row, column):
            cell = MagicMock()
            cell.value = row_values.get(row, 0.0)
            return cell

        ws.cell.side_effect = cell_side_effect
        hds._apply_profit_colors(ws, 3, 5, profit_col=12, rate_col=13, today_col=14)
        self.assertEqual(mock_pf.call_count, 9)

    def test_cell_font_assigned(self):
        """确保 cell.font 被赋值。"""
        ws = MagicMock()
        cell = MagicMock()
        cell.value = 100.0
        ws.cell.return_value = cell
        with patch("src.python.report.holdings_detail_sheet.profit_font") as mock_pf:
            mock_pf.return_value = "red_font"
            hds._apply_profit_colors(ws, 3, 3, profit_col=12, rate_col=13, today_col=14)
        self.assertEqual(cell.font, "red_font")

    def test_int_rate_value(self):
        """收益率列为 int → 不调用 profit_font（仅 float 触发）。"""
        with patch("src.python.report.holdings_detail_sheet.profit_font") as mock_pf:
            ws = MagicMock()
            # 仅 rate 列设为 int；profit/today 列 None（非数字）
            values = {13: 5}
            ws.cell.side_effect = lambda row, column: self._make_cell(values.get(column))
            hds._apply_profit_colors(ws, 3, 3, profit_col=12, rate_col=13, today_col=14)
            mock_pf.assert_not_called()

    def test_empty_range_no_crash(self):
        """空范围（start > end）→ 不报错，且不访问任何单元格。"""
        with patch("src.python.report.holdings_detail_sheet.profit_font"):
            ws = MagicMock()
            hds._apply_profit_colors(ws, 100, 50, profit_col=12, rate_col=13, today_col=14)
            ws.cell.assert_not_called()


# ═══════════════════════════════════════════════════════════
#  _apply_price_type_colors
# ═══════════════════════════════════════════════════════════


class TestApplyPriceTypeColors(unittest.TestCase):
    """测试 _apply_price_type_colors 取价方式列着色（使用真实 openpyxl Worksheet）。"""

    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.test_cases = [
            ("电池ETF", "场内收盘价(T)", True),
            ("长江电力", "场内收盘价(T-1)", False),
            ("中欧医疗", "官方净值(T)", True),
            ("某混合基金", "官方净值(T-1)", False),
            ("标普500ETF(QDII)", "官方净值(T-1)", True),
            ("恒生ETF(QDII)", "场内收盘价(T-1)", False),
            ("宁德时代", "场内实时价", False),
            ("海外收益(QDII)", "官方净值(T-2)", False),
            ("--", "--", False),
        ]
        for i, (name, price_type, expected_blue) in enumerate(self.test_cases):
            row = i + 2
            self.ws.cell(row=row, column=2, value=name)
            self.ws.cell(row=row, column=7, value=price_type)

    def _assert_blue(self, row: int, msg: str = ""):
        cell = self.ws.cell(row=row, column=7)
        self.assertIsNotNone(cell.font.color, f"Row {row} font.color is None")
        self.assertEqual(str(cell.font.color.rgb), "000066CC", msg)

    def _assert_not_blue(self, row: int, msg: str = ""):
        cell = self.ws.cell(row=row, column=7)
        if cell.font.color and cell.font.color.rgb:
            self.assertNotEqual(str(cell.font.color.rgb), "000066CC", msg)

    def test_scenario(self):
        """所有场景批量验证。"""
        hds._apply_price_type_colors(self.ws, 2, 2 + len(self.test_cases) - 1)
        errors = []
        for i, (name, price_type, expected_blue) in enumerate(self.test_cases):
            row = i + 2
            try:
                if expected_blue:
                    self._assert_blue(row, f"Row {row}: {name} / {price_type} should be blue")
                else:
                    self._assert_not_blue(row, f"Row {row}: {name} / {price_type} should NOT be blue")
            except AssertionError as e:
                errors.append(str(e))
        if errors:
            self.fail("\n".join(errors))

    def test_empty_range_no_error(self):
        """空范围（start > end）→ 不报错，且不访问任何单元格。"""
        ws = MagicMock()
        hds._apply_price_type_colors(ws, 100, 50)
        ws.cell.assert_not_called()

    def test_single_row(self):
        """单行范围。"""
        self.ws.cell(row=3, column=2, value="测试")
        self.ws.cell(row=3, column=7, value="场内收盘价(T)")
        hds._apply_price_type_colors(self.ws, 3, 3)
        self._assert_blue(3)

    def test_none_price_type_col(self):
        """取价方式列为 None → 不报错。"""
        row = 2 + len(self.test_cases)
        self.ws.cell(row=row, column=2, value="测试")
        self.ws.cell(row=row, column=7, value=None)
        hds._apply_price_type_colors(self.ws, row, row)
        self.assertIsNone(self.ws.cell(row=row, column=7).value)


# ═══════════════════════════════════════════════════════════
#  _write_account_groupings
# ═══════════════════════════════════════════════════════════


class TestWriteAccountGroupings(unittest.TestCase):
    """测试 _write_account_groupings 账户分组写入逻辑。"""

    def setUp(self):
        self.detail_a = hds.DetailRow(
            account="证券账户",
            name="电池ETF",
            code="561910",
            price=10.0,
            nav_date="2026-06-26",
            yesterday_close=9.5,
            price_type="场内收盘价(T)",
            premium="--",
            shares=100.0,
            market_value=1000.0,
            cost=100.0,
            profit=900.0,
            profit_rate=9.0,
            today_profit=50.0,
            source="腾讯财经",
            source_api="tencent",
        )
        self.detail_b = hds.DetailRow(
            account="支付宝",
            name="中欧医疗健康混合",
            code="003095",
            price=1.5,
            nav_date="2026-06-25",
            yesterday_close=1.48,
            price_type="官方净值(T-1)",
            premium="--",
            shares=200.0,
            market_value=300.0,
            cost=400.0,
            profit=-100.0,
            profit_rate=-0.25,
            today_profit=4.0,
            source="东方财富",
            source_api="eastmoney",
        )

    @patch("src.python.report.holdings_detail_sheet.write_data_row")
    @patch("src.python.report.holdings_detail_sheet.write_subtotal_row")
    @patch("src.python.report.holdings_detail_sheet._detail_to_row_values")
    @patch("src.python.report.holdings_detail_sheet._mv_num_formats")
    def test_single_account(self, mock_fmts, mock_to_row, mock_sub, mock_data):
        """单账户 → 1 个小计行。"""
        mock_fmts.return_value = [""] * 15
        mock_to_row.side_effect = lambda d: [d.account, d.name]
        ws = MagicMock()
        _, _, _, _, final_row = hds._write_account_groupings(ws, [self.detail_a], 3)
        mock_sub.assert_called_once()
        self.assertGreater(final_row, 3)

    @patch("src.python.report.holdings_detail_sheet.write_data_row")
    @patch("src.python.report.holdings_detail_sheet.write_subtotal_row")
    @patch("src.python.report.holdings_detail_sheet._detail_to_row_values")
    @patch("src.python.report.holdings_detail_sheet._mv_num_formats")
    def test_multiple_accounts(self, mock_fmts, mock_to_row, mock_sub, mock_data):
        """多账户 → 每个账户 1 个小计行。"""
        mock_fmts.return_value = [""] * 15
        mock_to_row.side_effect = lambda d: [d.account, d.name]
        ws = MagicMock()
        _, _, _, _, final_row = hds._write_account_groupings(ws, [self.detail_a, self.detail_b], 3)
        self.assertEqual(mock_sub.call_count, 2)

    @patch("src.python.report.holdings_detail_sheet.write_data_row")
    @patch("src.python.report.holdings_detail_sheet.write_subtotal_row")
    @patch("src.python.report.holdings_detail_sheet._detail_to_row_values")
    @patch("src.python.report.holdings_detail_sheet._mv_num_formats")
    def test_empty_details(self, mock_fmts, mock_to_row, mock_sub, mock_data):
        """空明细 → 无小计行。"""
        mock_fmts.return_value = [""] * 15
        ws = MagicMock()
        _, _, _, _, final_row = hds._write_account_groupings(ws, [], 3)
        mock_sub.assert_not_called()
        self.assertEqual(final_row, 3)

    @patch("src.python.report.holdings_detail_sheet.write_data_row")
    @patch("src.python.report.holdings_detail_sheet.write_subtotal_row")
    @patch("src.python.report.holdings_detail_sheet._detail_to_row_values")
    @patch("src.python.report.holdings_detail_sheet._mv_num_formats")
    def test_special_char_account(self, mock_fmts, mock_to_row, mock_sub, mock_data):
        """特殊字符账户名 → 不崩溃。"""
        mock_fmts.return_value = [""] * 15
        mock_to_row.side_effect = lambda d: [d.account, d.name]
        detail_special = hds.DetailRow(
            account="测💹试/账户（定投）",
            name="基金A",
            code="000001",
            shares=100.0,
            market_value=1000.0,
            cost=500.0,
            profit=500.0,
            profit_rate=1.0,
            today_profit=10.0,
        )
        ws = MagicMock()
        try:
            hds._write_account_groupings(ws, [detail_special], 3)
        except Exception as e:
            self.fail(f"特殊字符账户名崩溃: {e}")

    @patch("src.python.report.holdings_detail_sheet.write_data_row")
    @patch("src.python.report.holdings_detail_sheet.write_subtotal_row")
    @patch("src.python.report.holdings_detail_sheet._detail_to_row_values")
    @patch("src.python.report.holdings_detail_sheet._mv_num_formats")
    def test_acc_cost_zero_rate(self, mock_fmts, mock_to_row, mock_sub, mock_data):
        """成本为 0 时收益率为 0.0 而非除零异常。"""
        mock_fmts.return_value = [""] * 15
        mock_to_row.side_effect = lambda d: [d.account, d.name]
        detail_cost_zero = hds.DetailRow(
            account="证券",
            name="新股",
            code="688001",
            market_value=1000.0,
            cost=0.0,
            profit=1000.0,
            shares=100.0,
            profit_rate=None,
            today_profit=0.0,
        )
        ws = MagicMock()
        hds._write_account_groupings(ws, [detail_cost_zero], 3)
        sub_call = mock_sub.call_args
        if sub_call:
            subtotal_vals = sub_call[0][3]
            self.assertEqual(subtotal_vals[11], 0.0)


# ═══════════════════════════════════════════════════════════
#  _write_market_value_block（区块①）
# ═══════════════════════════════════════════════════════════


class TestWriteMarketValueBlock(unittest.TestCase):
    """测试 _write_market_value_block 区块①写入（mock 内部函数和 Excel 写入）。"""

    def setUp(self):
        self.holdings = [
            Holding("证券账户", "电池ETF", "561910", 100.0, 1.0),
            Holding("支付宝", "中欧医疗健康混合", "003095", 200.0, 2.0),
        ]
        self.details = [
            hds.DetailRow(
                account="证券账户",
                name="电池ETF",
                code="561910",
                price=10.0,
                nav_date="2026-06-26",
                yesterday_close=9.5,
                price_type="场内收盘价(T)",
                premium="--",
                shares=100.0,
                market_value=1000.0,
                cost=100.0,
                profit=900.0,
                profit_rate=9.0,
                today_profit=50.0,
                source="腾讯财经",
                source_api="tencent",
            ),
            hds.DetailRow(
                account="支付宝",
                name="中欧医疗健康混合",
                code="003095",
                price=1.5,
                nav_date="2026-06-25",
                yesterday_close=1.48,
                price_type="官方净值(T-1)",
                premium="--",
                shares=200.0,
                market_value=300.0,
                cost=400.0,
                profit=-100.0,
                profit_rate=-0.25,
                today_profit=4.0,
                source="东方财富",
                source_api="eastmoney",
            ),
        ]

    @patch("src.python.report.holdings_detail_sheet.write_total_row")
    @patch("src.python.report.holdings_detail_sheet.write_subtotal_row")
    @patch("src.python.report.holdings_detail_sheet.write_data_row")
    @patch("src.python.report.holdings_detail_sheet.write_header_row")
    @patch("src.python.report.holdings_detail_sheet.write_title_row")
    @patch("src.python.report.holdings_detail_sheet._apply_price_type_colors")
    @patch("src.python.report.holdings_detail_sheet._apply_profit_colors")
    @patch("src.python.report.holdings_detail_sheet.freeze_header")
    @patch("src.python.report.holdings_detail_sheet.auto_width")
    @patch("src.python.report.holdings_detail_sheet._detail_to_row_values")
    @patch("src.python.report.holdings_detail_sheet._mv_num_formats")
    def test_basic_write(
        self,
        mock_fmts,
        mock_to_row,
        mock_aw,
        mock_freeze,
        mock_color,
        mock_pt_color,
        mock_tl,
        mock_hdr,
        mock_data,
        mock_sub,
        mock_total,
    ):
        """正常写入：验证汇总值正确，内部函数被调用。"""
        mock_fmts.return_value = [""] * 15
        mock_to_row.side_effect = lambda d: [
            d.account,
            d.name,
            d.code,
            d.price,
            d.nav_date,
            d.yesterday_close,
            d.price_type,
            d.premium,
            d.shares,
            d.market_value,
            d.cost,
            d.profit,
            d.profit_rate,
            d.today_profit,
            d.source,
        ]
        mock_tl.return_value = 2
        mock_hdr.return_value = 3
        ws = MagicMock()
        ws.title = "fixture_title"
        result = hds._write_market_value_block(ws, self.details, None, 1)
        grand_mv, grand_cost, grand_profit, grand_today, details, _next_row = result
        self.assertAlmostEqual(grand_mv, 1300.0)
        self.assertAlmostEqual(grand_cost, 500.0)
        self.assertAlmostEqual(grand_profit, 800.0)
        self.assertAlmostEqual(grand_today, 54.0)
        self.assertEqual(len(details), 2)
        mock_tl.assert_called_once()
        mock_hdr.assert_called_once()
        self.assertEqual(mock_sub.call_count, 2)
        mock_total.assert_called_once()
        mock_color.assert_called_once()
        mock_pt_color.assert_called_once()
        # 冻结窗格由合并章节写入器负责（标题 + 区块小节标题 + 表头共 3 行），区块写入器不冻结
        mock_freeze.assert_not_called()
        # 列宽自适应同样由合并章节写入器统一执行
        mock_aw.assert_not_called()
        self.assertEqual(ws.title, "fixture_title")

    @patch("src.python.report.holdings_detail_sheet.write_total_row")
    @patch("src.python.report.holdings_detail_sheet.write_subtotal_row")
    @patch("src.python.report.holdings_detail_sheet.write_data_row")
    @patch("src.python.report.holdings_detail_sheet.write_header_row")
    @patch("src.python.report.holdings_detail_sheet.write_title_row")
    @patch("src.python.report.holdings_detail_sheet._apply_price_type_colors")
    @patch("src.python.report.holdings_detail_sheet._apply_profit_colors")
    @patch("src.python.report.holdings_detail_sheet.freeze_header")
    @patch("src.python.report.holdings_detail_sheet.auto_width")
    @patch("src.python.report.holdings_detail_sheet._detail_to_row_values")
    @patch("src.python.report.holdings_detail_sheet._mv_num_formats")
    def test_empty_holdings(
        self,
        mock_fmts,
        mock_to_row,
        mock_aw,
        mock_freeze,
        mock_color,
        mock_pt_color,
        mock_tl,
        mock_hdr,
        mock_data,
        mock_sub,
        mock_total,
    ):
        """空持仓 → 总市值为 0，无小计行。"""
        mock_fmts.return_value = [""] * 15
        mock_tl.return_value = 2
        mock_hdr.return_value = 3
        ws = MagicMock()
        result = hds._write_market_value_block(ws, [], None, 1)
        grand_mv, grand_cost, grand_profit, grand_today, details, _next_row = result
        self.assertAlmostEqual(grand_mv, 0.0)
        self.assertAlmostEqual(grand_cost, 0.0)
        self.assertAlmostEqual(grand_profit, 0.0)
        self.assertAlmostEqual(grand_today, 0.0)
        self.assertEqual(details, [])
        mock_sub.assert_not_called()
        mock_total.assert_called_once()

    @patch("src.python.report.holdings_detail_sheet.write_total_row")
    @patch("src.python.report.holdings_detail_sheet.write_subtotal_row")
    @patch("src.python.report.holdings_detail_sheet.write_data_row")
    @patch("src.python.report.holdings_detail_sheet.write_header_row")
    @patch("src.python.report.holdings_detail_sheet.write_title_row")
    @patch("src.python.report.holdings_detail_sheet._apply_price_type_colors")
    @patch("src.python.report.holdings_detail_sheet._apply_profit_colors")
    @patch("src.python.report.holdings_detail_sheet.freeze_header")
    @patch("src.python.report.holdings_detail_sheet.auto_width")
    @patch("src.python.report.holdings_detail_sheet._detail_to_row_values")
    @patch("src.python.report.holdings_detail_sheet._mv_num_formats")
    def test_subtotal_per_account(
        self,
        mock_fmts,
        mock_to_row,
        mock_aw,
        mock_freeze,
        mock_color,
        mock_pt_color,
        mock_tl,
        mock_hdr,
        mock_data,
        mock_sub,
        mock_total,
    ):
        """多个账户 → 每个账户写入小计。"""
        detail_a = self.details[0]
        detail_b = self.details[1]
        detail_c = hds.DetailRow(
            account="证券账户",
            name="长江电力",
            code="600900",
            price=25.0,
            nav_date="2026-06-26",
            yesterday_close=24.5,
            price_type="场内收盘价(T)",
            premium="--",
            shares=100.0,
            market_value=2500.0,
            cost=2000.0,
            profit=500.0,
            profit_rate=0.25,
            today_profit=50.0,
            source="腾讯财经",
            source_api="tencent",
        )
        mock_fmts.return_value = [""] * 15
        mock_to_row.side_effect = lambda d: [
            d.account,
            d.name,
            d.code,
            d.price,
            d.nav_date,
            d.yesterday_close,
            d.price_type,
            d.premium,
            d.shares,
            d.market_value,
            d.cost,
            d.profit,
            d.profit_rate,
            d.today_profit,
            d.source,
        ]
        mock_tl.return_value = 2
        mock_hdr.return_value = 3
        ws = MagicMock()
        result = hds._write_market_value_block(ws, [detail_a, detail_c, detail_b], None, 1)
        self.assertEqual(mock_sub.call_count, 2)
        grand_mv = result[0]
        self.assertAlmostEqual(grand_mv, 3800.0)

    @patch("src.python.report.holdings_detail_sheet.write_total_row")
    @patch("src.python.report.holdings_detail_sheet.write_subtotal_row")
    @patch("src.python.report.holdings_detail_sheet.write_data_row")
    @patch("src.python.report.holdings_detail_sheet.write_header_row")
    @patch("src.python.report.holdings_detail_sheet.write_title_row")
    @patch("src.python.report.holdings_detail_sheet._apply_price_type_colors")
    @patch("src.python.report.holdings_detail_sheet._apply_profit_colors")
    @patch("src.python.report.holdings_detail_sheet.freeze_header")
    @patch("src.python.report.holdings_detail_sheet.auto_width")
    @patch("src.python.report.holdings_detail_sheet._detail_to_row_values")
    @patch("src.python.report.holdings_detail_sheet._mv_num_formats")
    def test_all_zero_price_show_warning(
        self,
        mock_fmts,
        mock_to_row,
        mock_aw,
        mock_freeze,
        mock_color,
        mock_pt_color,
        mock_tl,
        mock_hdr,
        mock_data,
        mock_sub,
        mock_total,
    ):
        """全零行情 → 写入红色警告行 + 合并单元格。"""
        mock_fmts.return_value = [""] * 15
        mock_to_row.side_effect = lambda d: [""] * 15
        mock_tl.return_value = 2
        mock_hdr.return_value = 3
        ws = MagicMock()
        ws.cell.return_value = MagicMock()
        hds._write_market_value_block(
            ws,
            [
                hds.DetailRow(
                    account="证券", name="电池ETF", code="561910", price=0.0, shares=100.0, market_value=0.0, cost=100.0
                ),
            ],
            None,
            1,
        )
        ws.merge_cells.assert_called_once()
        call_kwargs = ws.cell.call_args
        if call_kwargs:
            written_cell = ws.cell.return_value
            self.assertIsNotNone(written_cell.font)


# ═══════════════════════════════════════════════════════════
#  _weighted_avg_cost — 资金加权成本计算（成本流水子模块）
# ═══════════════════════════════════════════════════════════


class TestWeightedAvgCost(unittest.TestCase):
    """_weighted_avg_cost 单元测试：批次成本价按份额加权（含费用摊薄）。"""

    def test_returns_weighted_cost(self):
        """多档批次按份额加权 = 总成本 / 总份额。"""
        buckets = {
            "low": {"shares": 100.0, "cost": 900.0},
            "high": {"shares": 50.0, "cost": 600.0},
            "unpriced": {"shares": 0.0, "cost": 0.0},
        }
        # (900 + 600) / (100 + 50) = 1500 / 150 = 10.0
        self.assertAlmostEqual(hds._weighted_avg_cost(buckets), 10.0)

    def test_single_low_bucket(self):
        """仅低成本档时返回其平均成本价。"""
        buckets = {
            "low": {"shares": 100.0, "cost": 950.0},
            "high": {"shares": 0.0, "cost": 0.0},
            "unpriced": {"shares": 0.0, "cost": 0.0},
        }
        self.assertAlmostEqual(hds._weighted_avg_cost(buckets), 9.5)

    def test_none_for_missing_or_empty(self):
        """缺码/空桶/零份额时返回 None（渲染层写空）。"""
        self.assertIsNone(hds._weighted_avg_cost(None))
        self.assertIsNone(hds._weighted_avg_cost({}))
        self.assertIsNone(hds._weighted_avg_cost({}))
        self.assertIsNone(
            hds._weighted_avg_cost({"low": {"shares": 0.0, "cost": 0.0}, "high": {"shares": 0.0, "cost": 0.0}})
        )


# ═══════════════════════════════════════════════════════════
#  资金加权成本列渲染（成本流水子模块）
# ═══════════════════════════════════════════════════════════


class TestWriteMarketValueBlockFlow(unittest.TestCase):
    """_write_market_value_block 资金加权成本列渲染测试。

    功能开关 `cost_lots` 对应 fund_flow_data 是否传入：
    None → 保持既有 15 列；非 None → 追加「资金加权成本」列（第 16 列）。
    """

    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.detail = hds.DetailRow(
            account="证券账户",
            name="电池ETF",
            code="561910",
            price=10.0,
            nav_date="2026-06-26",
            yesterday_close=9.5,
            price_type="场内收盘价(T)",
            premium="--",
            shares=100.0,
            market_value=1000.0,
            cost=100.0,
            profit=900.0,
            profit_rate=9.0,
            today_profit=50.0,
            source="腾讯财经",
            source_api="tencent",
        )

    def _flow(self, low_shares=0.0, cost=0.0):
        return {
            "available": True,
            "cost_tiers": {
                "per_code": {
                    "561910": {
                        "low": {"shares": low_shares, "cost": cost},
                        "high": {"shares": 0.0, "cost": 0.0},
                        "unpriced": {"shares": 0.0, "cost": 0.0},
                    }
                }
            },
            "dividends": {"per_code": {}},
        }

    def test_flow_weighted_cost_header_when_enabled(self):
        """开关开启时表头第 16 列为「资金加权成本」，原 15 列保持不变。"""
        hds._write_market_value_block(self.ws, [self.detail], self._flow(), 1)
        headers = [self.ws.cell(row=2, column=c).value for c in range(1, 17)]
        self.assertEqual(len(headers), 16)
        self.assertEqual(headers[14], "取价渠道")
        self.assertEqual(headers[15], "资金加权成本")

    def test_flow_weighted_cost_value(self):
        """开关开启时数据行「资金加权成本」列 = 批次成本价按份额加权。"""
        flow = self._flow(low_shares=100.0, cost=950.0)
        hds._write_market_value_block(self.ws, [self.detail], flow, 1)
        # 数据行 row 3，列 16 = 950 / 100 = 9.5
        self.assertAlmostEqual(self.ws.cell(row=3, column=16).value, 9.5)

    def test_flow_weighted_cost_none_when_no_buckets(self):
        """开关开启但代码无批次数据时，「资金加权成本」列为空。"""
        flow = {"available": True, "cost_tiers": {"per_code": {}}, "dividends": {"per_code": {}}}
        hds._write_market_value_block(self.ws, [self.detail], flow, 1)
        self.assertIsNone(self.ws.cell(row=3, column=16).value)

    def test_no_flow_column_when_disabled(self):
        """开关关闭（fund_flow_data=None）时保持既有 15 列，无「资金加权成本」列。"""
        hds._write_market_value_block(self.ws, [self.detail], None, 1)
        headers = [self.ws.cell(row=2, column=c).value for c in range(1, 16)]
        self.assertEqual(len(headers), 15)
        self.assertEqual(headers[14], "取价渠道")
        self.assertNotIn("资金加权成本", headers)


# ═══════════════════════════════════════════════════════════
#  区块② 持仓分类汇总（由 test_category.py 迁入）
# ═══════════════════════════════════════════════════════════


class TestWriteCategorySheet(unittest.TestCase):
    """区块② 持仓分类汇总写入集成测试。"""

    def setUp(self):
        import openpyxl

        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active

    def test_single_holding(self):
        """单条持仓 → 分类表正确渲染。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(code="600900", account="证券", name="长江电力", market_value=2500, profit=1500, profit_rate=1.5)
        ]
        hds._write_category_block(self.ws, [h], details, None, 1)
        # 应该至少有一行标题 + 一行数据 + 合计行
        self.assertGreater(self.ws.max_row, 2)

    def test_empty_holdings(self):
        """空持仓 → 不崩溃。"""
        try:
            hds._write_category_block(self.ws, [], [], None, 1)
        except Exception as e:
            self.fail(f"空持仓分类区块写入不应抛异常: {e}")


class TestWriteCategoryFlowSubcolumns(unittest.TestCase):
    """持仓分类表成本流水子列（成本分档 / 分红累计）渲染测试。

    开关 `report_submodules.cost_lots` 对应 fund_flow_data 是否传入：
    None → 保持既有 10 列；非 None → 追加「成本分档」「分红累计」子列。
    """

    def setUp(self):
        import openpyxl

        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active

    def _call(self, holdings, details, fund_flow_data=None):
        # mock 分红 API 加载（避免真实网络请求，测试隔离）
        with patch("src.python.report.category._load_dividend_data", return_value=({}, True)):
            hds._write_category_block(self.ws, holdings, details, fund_flow_data, 1)

    def _headers(self):
        return [self.ws.cell(row=2, column=c).value for c in range(1, 13)]

    def _find_total_row(self):
        for r in range(1, self.ws.max_row + 1):
            if self.ws.cell(row=r, column=1).value == "总计":
                return r
        return None

    @staticmethod
    def _flow(low_shares=0.0, high_shares=0.0, div=0.0):
        return {
            "available": True,
            "cost_tiers": {
                "per_code": {
                    "600900": {
                        "low": {"shares": low_shares, "cost": low_shares * 9.0},
                        "high": {"shares": high_shares, "cost": high_shares * 11.0},
                        "unpriced": {"shares": 0.0, "cost": 0.0},
                    }
                }
            },
            "dividends": {"per_code": {"600900": div}},
        }

    def test_flow_subcolumns_header_when_enabled(self):
        """开关开启（fund_flow_data 非 None）时，表头追加「成本分档」「分红累计」列。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=1500,
                cost=1000,
                profit=500,
                profit_rate=0.5,
            )
        ]
        self._call([h], details, self._flow(low_shares=100, div=120.0))
        headers = self._headers()
        self.assertIn("成本分档", headers)
        self.assertIn("分红累计", headers)
        self.assertEqual(headers.index("成本分档") + 1, 11)
        self.assertEqual(headers.index("分红累计") + 1, 12)

    def test_flow_tier_low_and_div_value(self):
        """开关开启时数据行含「低成本」分档标签与分红累计数值。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=1500,
                cost=1000,
                profit=500,
                profit_rate=0.5,
            )
        ]
        self._call([h], details, self._flow(low_shares=100, div=120.0))
        # 数据行 = 表头下一行（row 3）：列 11 = 成本分档, 列 12 = 分红累计
        self.assertEqual(self.ws.cell(row=3, column=11).value, "低成本")
        self.assertEqual(self.ws.cell(row=3, column=12).value, 120.0)

    def test_flow_tier_high_label(self):
        """持仓批次成本价高于市价时渲染「高成本」档标签。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=900,
                cost=1000,
                profit=-100,
                profit_rate=-0.1,
            )
        ]
        self._call([h], details, self._flow(high_shares=100, div=0.0))
        self.assertEqual(self.ws.cell(row=3, column=11).value, "高成本")

    def test_flow_tier_mixed_label(self):
        """持仓批次横跨低/高两档时渲染「混合」档标签。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900", account="证券", name="长江电力", market_value=1000, cost=1000, profit=0, profit_rate=0.0
            )
        ]
        self._call([h], details, self._flow(low_shares=50, high_shares=50, div=0.0))
        self.assertEqual(self.ws.cell(row=3, column=11).value, "混合")

    def test_flow_total_row_dividend_sum(self):
        """开关开启时，分红累计在小计/总计行汇总（跨持仓累加）。"""
        flow = {
            "available": True,
            "cost_tiers": {
                "per_code": {
                    "600900": {
                        "low": {"shares": 100, "cost": 900},
                        "high": {"shares": 0, "cost": 0},
                        "unpriced": {"shares": 0, "cost": 0},
                    },
                    "600519": {
                        "low": {"shares": 10, "cost": 15000},
                        "high": {"shares": 0, "cost": 0},
                        "unpriced": {"shares": 0, "cost": 0},
                    },
                }
            },
            "dividends": {"per_code": {"600900": 120.0, "600519": 30.0}},
        }
        h1 = Holding("证券", "长江电力", "600900", 100, 10.0)
        h2 = Holding("证券", "贵州茅台", "600519", 10, 1500.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=1500,
                cost=1000,
                profit=500,
                profit_rate=0.5,
            ),
            DetailRow(
                code="600519",
                account="证券",
                name="贵州茅台",
                market_value=20000,
                cost=15000,
                profit=5000,
                profit_rate=0.33,
            ),
        ]
        self._call([h1, h2], details, flow)
        total_row = self._find_total_row()
        self.assertIsNotNone(total_row)
        self.assertEqual(self.ws.cell(row=total_row, column=12).value, 150.0)

    def test_no_flow_subcolumns_when_disabled(self):
        """开关关闭（fund_flow_data=None）时保持既有 10 列输出，无成本流水子列。"""
        h = Holding("证券", "长江电力", "600900", 100, 10.0)
        details = [
            DetailRow(
                code="600900",
                account="证券",
                name="长江电力",
                market_value=1500,
                cost=1000,
                profit=500,
                profit_rate=0.5,
            )
        ]
        self._call([h], details, None)
        headers = self._headers()
        self.assertNotIn("成本分档", headers)
        self.assertNotIn("分红累计", headers)
        # 既有表头 10 列保持不变（前 10 列与改造前一致）
        self.assertEqual(headers[:10][-1], "年均股息率")


@pytest.mark.data

# ═══════════════════════════════════════════════════════════
#  write_holdings_detail_sheet — 合并章节写入器（区块① + 区块②）
# ═══════════════════════════════════════════════════════════


class TestWriteHoldingsDetailSheet(unittest.TestCase):
    """合并章节写入器：单页签承载两个区块（持仓明细与分类）。"""

    def setUp(self):
        import openpyxl

        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.details = [
            hds.DetailRow(
                account="A",
                name="贵州茅台",
                code="600519",
                price=1500.0,
                nav_date="",
                yesterday_close=1480.0,
                price_type="场内收盘价(T)",
                premium="--",
                shares=100.0,
                market_value=150000.0,
                cost=140000.0,
                profit=10000.0,
                profit_rate=0.0714,
                today_profit=2000.0,
                source="腾讯财经",
            ),
            hds.DetailRow(
                account="A",
                name="某基金",
                code="040046",
                price=1.234,
                nav_date="2026-09-14",
                yesterday_close=1.22,
                price_type="官方净值(T-1)",
                premium="--",
                shares=1000.0,
                market_value=1234.0,
                cost=1200.0,
                profit=34.0,
                profit_rate=0.028,
                today_profit=14.0,
                source="基金净值",
            ),
        ]
        self.holdings = [
            Holding(code="600519", name="贵州茅台", shares=100.0, cost_price=1400.0, account="A"),
            Holding(code="040046", name="某基金", shares=1000.0, cost_price=1.2, account="A"),
        ]

    def test_title_is_merged_sheet_name(self):
        """首行为合并章节显示名「持仓明细与分类」。"""
        from src.python.core.registry import get_report_sheet_name

        hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details)
        self.assertEqual(self.ws["A1"].value, get_report_sheet_name("holdings_detail"))
        self.assertEqual(self.ws["A1"].value, "持仓明细与分类")

    def test_two_blocks_present_in_one_sheet(self):
        """同一工作表内含两个区块小节标题（市值明细 / 分类汇总）。"""
        hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details)
        col_a = [self.ws.cell(row=r, column=1).value for r in range(1, self.ws.max_row + 1)]
        self.assertIn("一、市值核算明细", col_a)
        self.assertIn("二、持仓分类汇总", col_a)
        # 区块①表头与区块②表头各自出现（同页签两表）
        all_vals = [self.ws.cell(row=r, column=c).value for r in range(1, self.ws.max_row + 1) for c in range(1, 11)]
        self.assertIn("账户", all_vals)
        self.assertIn("资产属性", all_vals)

    def test_returns_market_summary_tuple(self):
        """返回值与原市值明细写入器一致：(总市值, 总成本, 总盈亏, 本日总盈亏, 明细行)。"""
        result = hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details)
        self.assertEqual(len(result), 5)
        self.assertAlmostEqual(result[0], 151234.0)
        self.assertAlmostEqual(result[1], 141200.0)
        self.assertAlmostEqual(result[2], 10034.0)
        self.assertAlmostEqual(result[3], 2014.0)
        self.assertEqual(len(result[4]), 2)

    def test_block_values_equal_standalone_blocks(self):
        """内容等价：合并页签的两个区块行值 == 各自独立写入（start_row=1）的结果。"""
        import openpyxl

        hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details)

        wb2 = openpyxl.Workbook()
        ws_mv = wb2.active
        hds._write_market_value_block(ws_mv, self.details, None, 1)
        ws_cat = wb2.create_sheet()
        hds._write_category_block(ws_cat, self.holdings, self.details, None, 1)

        # 区块①：跳过合并页签新增的区块小节标题行（第 2 行），其余逐行相等
        mv_merged = [
            [self.ws.cell(row=r + 1, column=c).value for c in range(1, 17)] for r in range(1, ws_mv.max_row + 1)
        ]
        mv_standalone = [[ws_mv.cell(row=r, column=c).value for c in range(1, 17)] for r in range(1, ws_mv.max_row + 1)]
        self.assertEqual(mv_merged[2:], mv_standalone[2:])  # 表头及数据行逐格相等

        # 区块②：分类行数据（资产属性/投资分类/名称/代码/市值）与独立写入一致
        cat_merged = [
            [self.ws.cell(row=r, column=c).value for c in range(1, 11)]
            for r in range(1, self.ws.max_row + 1)
            if self.ws.cell(row=r, column=1).value in ("股票", "基金", "现金", "债券", "其他")
        ]
        cat_standalone = [
            [ws_cat.cell(row=r, column=c).value for c in range(1, 11)]
            for r in range(1, ws_cat.max_row + 1)
            if ws_cat.cell(row=r, column=1).value in ("股票", "基金", "现金", "债券", "其他")
        ]
        self.assertEqual(cat_merged, cat_standalone)
        self.assertTrue(cat_merged, "分类区块应至少有一行分类明细")

    def test_flow_columns_apply_to_both_blocks(self):
        """成本流水开关开启时，两个区块各自追加子列（区块①「资金加权成本」/ 区块②「成本分档」）。"""
        flow = {"available": True, "approximate": False, "cost_tiers": {"per_code": {}}, "dividends": {"per_code": {}}}
        hds.write_holdings_detail_sheet(self.ws, self.holdings, self.details, fund_flow_data=flow)
        all_vals = [self.ws.cell(row=r, column=c).value for r in range(1, self.ws.max_row + 1) for c in range(1, 17)]
        self.assertIn("资金加权成本", all_vals)
        self.assertIn("成本分档", all_vals)
        self.assertIn("分红累计", all_vals)


class TestAllZeroPriceRegression(unittest.TestCase):
    """缺陷回归：行情全零时提示行**整行合并**，数据必须自其后一行写入。

    现场（logs/app.log 2026-09-16 场景测试暴露）：提示行 `merge_cells` 后仍以该行为
    数据起点 → `AttributeError: 'MergedCell' object attribute 'value' is read-only`，
    导致整份 Excel 报告生成失败。既有实现（合并前的 market_value_sheet）同样有此缺陷，
    批次② 合并时原样带入。
    """

    def _ws(self):
        import openpyxl

        wb = openpyxl.Workbook()
        return wb.active

    def test_all_zero_price_writes_warning_and_data_without_crash(self):
        ws = self._ws()
        zero_details = [
            hds.DetailRow(
                account="A", name="贵州茅台", code="600519", price=0.0, shares=100.0, market_value=0.0, cost=140000.0
            ),
            hds.DetailRow(
                account="A", name="某基金", code="040046", price=0.0, shares=1000.0, market_value=0.0, cost=1200.0
            ),
        ]
        holdings = [
            Holding(code="600519", name="贵州茅台", shares=100.0, cost_price=1400.0, account="A"),
            Holding(code="040046", name="某基金", shares=1000.0, cost_price=1.2, account="A"),
        ]
        # 不得抛异常（缺陷现场在此崩溃）
        hds.write_holdings_detail_sheet(ws, holdings, zero_details)

        flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        self.assertTrue(any("行情数据全部不可用" in v for v in flat), "应写入全零提示行")
        self.assertIn("贵州茅台", flat, "数据行应写在提示行之后")
        self.assertIn("A 小计", flat)

    def test_all_zero_price_indicator_block_order(self):
        """区块① 中提示行之后的数据行号必须大于提示行行号（防写回合并区）。"""
        ws = self._ws()
        zero_details = [
            hds.DetailRow(
                account="A", name="贵州茅台", code="600519", price=0.0, shares=100.0, market_value=0.0, cost=140000.0
            )
        ]
        hds.write_holdings_detail_sheet(ws, [], zero_details)
        warn_row = next(
            r for r in range(1, ws.max_row + 1) if "行情数据全部不可用" in str(ws.cell(row=r, column=1).value or "")
        )
        data_row = next((r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "贵州茅台"), None)
        self.assertIsNotNone(data_row, "全零场景仍应写出数据行")
        self.assertGreater(data_row, warn_row, "数据行必须在提示行之后")
