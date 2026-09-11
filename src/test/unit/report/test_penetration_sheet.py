"""穿透页签写入层单元测试。

测试目标：
  - _write_penetration_footer — 底部备注的剔除原因、明细与各基金持仓报告期

运行：
  pytest src/test/unit/report/test_penetration_sheet.py -v
"""

from __future__ import annotations

import unittest

from openpyxl import Workbook

from src.python.report import penetration_sheet as ps
import pytest

pytestmark = [pytest.mark.unit, pytest.mark.unit_report]


def _sheet_text(ws) -> str:
    """把页签所有已写入的单元格拼成一段文本，便于断言备注内容。"""
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def _summary(**overrides) -> dict:
    """构造穿透 summary 最小集合，测试按需覆写。"""
    summary = {
        "total_funds": 2,
        "total_stocks": 0,
        "fund_breakdown": "主动2",
        "merged_count": 3,
        "top10_coverage_pct": 88.8,
        "unknown_mv": 0.0,
        "failed_funds": 0,
        "failed_fund_details": [],
        "stale_funds": 0,
        "stale_fund_details": [],
        "report_periods": [],
    }
    summary.update(overrides)
    return summary


class TestPenetrationFooterReportPeriods(unittest.TestCase):
    """底部备注须标出各基金持仓的报告期，否则穿透权重无从判断时点。"""

    def test_report_periods_are_listed(self):
        """各基金报告期逐只列出（名称 + 代码 + 报告期）。"""
        wb = Workbook()
        ws = wb.active
        summary = _summary(
            report_periods=[
                {"name": "易方达蓝筹", "code": "005827", "period": "2026-06-30"},
                {"name": "广发科技先锋", "code": "005911", "period": "2026-06-30"},
            ]
        )

        ps._write_penetration_footer(ws, 0, summary)

        text = _sheet_text(ws)
        self.assertIn("各基金持仓报告期", text)
        self.assertIn("易方达蓝筹(005827) 2026-06-30", text)
        self.assertIn("广发科技先锋(005911) 2026-06-30", text)

    def test_no_report_periods_no_extra_line(self):
        """无报告期明细时不写这一行（空框架不出现）。"""
        wb = Workbook()
        ws = wb.active

        ps._write_penetration_footer(ws, 0, _summary())

        self.assertNotIn("各基金持仓报告期", _sheet_text(ws))

    def test_feeder_fund_shows_penetration_source(self):
        """联接基金的持仓穿透自目标 ETF，须标出来源并明示未折算持有比例。

        未折算说明不可省略：联接基金约 95% 资产投向目标 ETF，按 100% 归因
        会轻微高估底层标的权重，读者当作精确值会误判真实暴露。
        """
        wb = Workbook()
        ws = wb.active
        summary = _summary(
            report_periods=[
                {
                    "name": "博时纳斯达克100ETF发起式联接(QDII)A人民币",
                    "code": "016055",
                    "period": "2026-06-30",
                    "feeder_target_code": "513390",
                    "feeder_target_name": "纳指100ETF博时",
                }
            ]
        )

        ps._write_penetration_footer(ws, 0, summary)

        text = _sheet_text(ws)
        self.assertIn("016055)", text)
        self.assertIn("穿透自目标 ETF 513390 纳指100ETF博时", text)
        self.assertIn("未折算持有比例", text)

    def test_non_feeder_fund_has_no_source_annotation(self):
        """非联接基金不带来源标注（避免给普通基金加上误导性说明）。"""
        wb = Workbook()
        ws = wb.active
        summary = _summary(report_periods=[{"name": "易方达蓝筹", "code": "005827", "period": "2026-06-30"}])

        ps._write_penetration_footer(ws, 0, summary)

        text = _sheet_text(ws)
        self.assertIn("易方达蓝筹(005827) 2026-06-30", text)
        self.assertNotIn("穿透自目标 ETF", text)


class TestPenetrationFooterStaleFunds(unittest.TestCase):
    """报告期陈旧被剔除的基金须在备注中说明原因与原始报告期。"""

    def test_stale_fund_reason_and_detail(self):
        """陈旧剔除单列计数与明细，不明说成「无法获取」。"""
        wb = Workbook()
        ws = wb.active
        summary = _summary(
            unknown_mv=1000.0,
            stale_funds=1,
            stale_fund_details=[
                {
                    "name": "华安纳斯达克100ETF联接(QDII)A",
                    "code": "040046",
                    "period": "2022-12-08",
                    "quarters": "14",
                }
            ],
        )

        ps._write_penetration_footer(ws, 0, summary)

        text = _sheet_text(ws)
        self.assertIn("1 只因持仓报告期陈旧被剔除", text)
        self.assertIn("合计市值 1,000.00 元未计入穿透 TOP10", text)
        self.assertIn("报告期陈旧被剔除的基金", text)
        self.assertIn("华安纳斯达克100ETF联接(QDII)A(040046) 报告期 2022-12-08，距今 14 个完整季度", text)

    def test_both_reasons_are_joined(self):
        """获取失败与陈旧剔除并存时，两种原因并列在同一行。"""
        wb = Workbook()
        ws = wb.active
        summary = _summary(
            unknown_mv=3000.0,
            failed_funds=1,
            failed_fund_details=[{"name": "新发基金", "code": "019999"}],
            stale_funds=1,
            stale_fund_details=[{"name": "陈旧基金", "code": "040046", "period": "2022-12-08", "quarters": "14"}],
        )

        ps._write_penetration_footer(ws, 0, summary)

        text = _sheet_text(ws)
        self.assertIn("1 只无法获取穿透数据、1 只因持仓报告期陈旧被剔除", text)

    def test_failed_only_keeps_original_wording(self):
        """仅获取失败时措辞与既有输出一致（不引入陈旧相关文案）。"""
        wb = Workbook()
        ws = wb.active
        summary = _summary(
            unknown_mv=500.0,
            failed_funds=1,
            failed_fund_details=[{"name": "新发基金", "code": "019999"}],
        )

        ps._write_penetration_footer(ws, 0, summary)

        text = _sheet_text(ws)
        self.assertIn("1 只无法获取穿透数据，", text)
        self.assertNotIn("报告期陈旧", text)


if __name__ == "__main__":
    unittest.main()
