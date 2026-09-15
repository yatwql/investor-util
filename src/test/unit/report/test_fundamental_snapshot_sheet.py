"""财务指标页签 Excel 呈现测试。

覆盖：available=True → 标题 + 表头 + 数值格式（亿元/百分数/倍数）+ 失败清单；
available=False → 占位原因；None → 整页占位；缺失值写「—」而非 0。

运行：
  pytest src/test/unit/report/test_financial_indicator_sheet.py -v
"""

from __future__ import annotations

import unittest

import pytest

pytestmark = [pytest.mark.unit, pytest.mark.unit_report]


def _indicator_data(**extra) -> dict:
    d = {
        "available": True,
        "reason": "",
        "entry_count": 1,
        "rows": [
            {
                "code": "600900",
                "name": "长江电力",
                "report_period": "2025-12-31",
                "doc_type": "annual",
                "doc_type_label": "年报",
                "source_api": "akshare_financial",
                "source": "akshare 财务指标",
                "revenue": 86_241_940_222.20,
                "net_profit": 34_502_809_176.39,
                "revenue_yoy": 0.0207,
                "net_profit_yoy": 0.0617,
                "gross_margin": 0.35,
                "roe": 0.159,
                "debt_ratio": 0.5827,
                "operating_cash_flow": 60_562_925_570.41,
                "eps": 1.4101,
                "bvps": 9.05,
                "pe": 20.21,
                "pb": 3.15,
                "quality_grade": "优",
                "quality_score": 2.75,
                "trend": "增长",
                "series": [],
                "period_count": 2,
            }
        ],
        "failures": [],
    }
    d.update(extra)
    return d


def _digest_data(**extra) -> dict:
    """构造 financial_report_digest_data 契约 mock（单只 A 股）。"""
    d = {
        "available": True,
        "reason": "",
        "entry_count": 1,
        "rows": [
            {
                "code": "600519",
                "name": "贵州茅台",
                "symbol": "600519.SS",
                "report_period": "2025-12-31",
                "doc_type": "年报",
                "title": "2025 年年度报告",
                "announcement_date": "2026-04-01",
                "summary": "管理层讨论与分析：营收同比增长……",
                "source": "巨潮资讯网 (cninfo)",
                "adjunct_url": "http://example.com/1",
            }
        ],
        "failures": [],
    }
    d.update(extra)
    return d


class TestIndicatorBlock(unittest.TestCase):
    def _write(self, indicator_data):
        from openpyxl import Workbook

        from src.python.report.fundamental_snapshot_sheet import _write_indicator_block

        wb = Workbook()
        ws = wb.active
        _write_indicator_block(ws, 1, indicator_data)
        return ws

    def _flat(self, ws) -> list[str]:
        return [str(c.value) if c.value is not None else "" for row in ws.iter_rows() for c in row]

    def test_available_writes_header_and_formatted_values(self):
        ws = self._write(_indicator_data())
        flat = self._flat(ws)
        self.assertTrue(any("财务指标" in v for v in flat))
        for expected in ("长江电力", "600900", "2025-12-31", "年报", "862.42", "345.03", "605.63"):
            self.assertTrue(any(expected in v for v in flat), f"缺少 {expected}: {flat}")
        # 比率换算为百分数；倍数保留两位
        self.assertTrue(any("6.17%" in v for v in flat), flat)
        self.assertTrue(any("15.90%" in v for v in flat), flat)
        self.assertTrue(any("20.21" in v for v in flat), flat)
        self.assertTrue(any("优" in v for v in flat), flat)

    def test_missing_values_render_dash_not_zero(self):
        row = dict(_indicator_data()["rows"][0])
        row.update({"gross_margin": None, "pe": None, "quality_grade": "", "trend": ""})
        ws = self._write(_indicator_data(rows=[row]))
        flat = self._flat(ws)
        self.assertIn("—", flat, "缺失值应写「—」而非 0")

    def test_unavailable_writes_reason_and_failures(self):
        ws = self._write(
            _indicator_data(
                available=False,
                reason="未取到财务指标数据（数据源不可用或标的不在覆盖范围）",
                rows=[],
                entry_count=0,
                failures=[{"code": "600900", "name": "长江电力", "reason": "无指标数据"}],
            )
        )
        flat = self._flat(ws)
        self.assertTrue(any("数据源不可用" in v for v in flat))
        self.assertTrue(any("无指标数据" in v for v in flat))

    def test_none_writes_placeholder(self):
        ws = self._write(None)
        flat = self._flat(ws)
        self.assertTrue(any("暂无可用财务指标数据" in v for v in flat))


# ═══════════════════════════════════════════════════════════
#  区块② 持仓个股财报摘要（由 test_financial_report_sheet.py 迁入）
# ═══════════════════════════════════════════════════════════


class TestDigestBlock(unittest.TestCase):
    def _write(self, digest_data):
        from openpyxl import Workbook

        from src.python.report.fundamental_snapshot_sheet import _write_digest_block

        wb = Workbook()
        ws = wb.active
        _write_digest_block(ws, 1, digest_data)
        return ws

    def _flat(self, ws) -> list[str]:
        return [str(c.value) if c.value is not None else "" for row in ws.iter_rows() for c in row]

    def test_available_writes_rows_and_notes(self):
        ws = self._write(_digest_data())
        flat = self._flat(ws)
        self.assertTrue(any("持仓个股财报摘要" in v for v in flat))
        for expected in ("贵州茅台", "600519", "2025-12-31", "年报", "2026-04-01", "巨潮资讯网", "原文"):
            self.assertTrue(any(expected in v for v in flat), f"缺少 {expected}: {flat}")

    def test_unavailable_writes_reason_and_failures(self):
        ws = self._write(
            _digest_data(
                available=False,
                reason="未配置 DataSinking API key",
                failures=[{"code": "600519", "name": "贵州茅台", "reason": "未取到财报"}],
            )
        )
        flat = self._flat(ws)
        self.assertTrue(any("未配置 DataSinking API key" in v for v in flat))
        self.assertTrue(any("未取到财报" in v for v in flat))

    def test_none_writes_placeholder(self):
        ws = self._write(None)
        flat = self._flat(ws)
        self.assertTrue(any("暂无可用财报数据" in v for v in flat))


# ═══════════════════════════════════════════════════════════
#  合并章节写入器：两区块 + 区块级开关门控
# ═══════════════════════════════════════════════════════════


class TestWriteFundamentalSnapshotSheet(unittest.TestCase):
    """合并章节：单页签承载两区块，契约 None（开关关闭）时该区块整体不写。"""

    def _write(self, indicator_data=None, digest_data=None):
        from openpyxl import Workbook

        from src.python.report.fundamental_snapshot_sheet import write_fundamental_snapshot_sheet

        wb = Workbook()
        ws = wb.active
        write_fundamental_snapshot_sheet(ws, indicator_data=indicator_data, digest_data=digest_data)
        return ws

    def _flat(self, ws) -> list[str]:
        return [str(c.value) if c.value is not None else "" for row in ws.iter_rows() for c in row]

    def test_title_and_both_blocks(self):
        """首行为合并章节显示名，两个区块小节标题同页签。"""
        from src.python.core.registry import get_report_sheet_name

        ws = self._write(_indicator_data(), _digest_data())
        flat = self._flat(ws)
        self.assertEqual(ws["A1"].value, get_report_sheet_name("fundamental_snapshot"))
        self.assertEqual(ws["A1"].value, "持仓基本面")
        self.assertTrue(any("一、财务指标" in v for v in flat), flat[:20])
        self.assertTrue(any("二、持仓个股财报摘要" in v for v in flat), flat[:20])

    def test_digest_block_hidden_when_contract_none(self):
        """财报摘要开关关闭（契约 None）→ 区块②整体不写（连小节标题也不出现）。"""
        ws = self._write(_indicator_data(), None)
        flat = self._flat(ws)
        self.assertTrue(any("一、财务指标" in v for v in flat))
        self.assertFalse(any("二、持仓个股财报摘要" in v for v in flat), "关闭的区块不得留标题")

    def test_indicator_block_hidden_when_contract_none(self):
        """财务指标开关关闭（契约 None）→ 区块①整体不写。"""
        ws = self._write(None, _digest_data())
        flat = self._flat(ws)
        self.assertFalse(any("一、财务指标" in v for v in flat), "关闭的区块不得留标题")
        self.assertTrue(any("二、持仓个股财报摘要" in v for v in flat))

    def test_both_contracts_absent_writes_no_block(self):
        """两开关均关闭 → 仅有章节标题（页签由 data_flag_any 悲观判定通常不创建）。"""
        ws = self._write(None, None)
        flat = [v for v in self._flat(ws) if v]
        self.assertEqual(flat, ["持仓基本面"])

    def test_block_values_equal_standalone_blocks(self):
        """内容等价：两区块的单元格矩阵与独立区块写入（起始行=1）一致。"""
        from openpyxl import Workbook

        from src.python.report.fundamental_snapshot_sheet import _write_digest_block, _write_indicator_block

        ws = self._write(_indicator_data(), _digest_data())
        wb2 = Workbook()
        ws_ind = wb2.active
        _write_indicator_block(ws_ind, 1, _indicator_data())
        ws_dig = wb2.create_sheet()
        _write_digest_block(ws_dig, 1, _digest_data())

        # 合并页签中，区块① 自「一、财务指标」起、区块② 自「二、持仓个股财报摘要」起
        def dump(sheet, marker, limit=19):
            rows, started = [], False
            for r in range(1, sheet.max_row + 1):
                first = sheet.cell(row=r, column=1).value
                if first == marker:
                    started = True
                if started:
                    rows.append([sheet.cell(row=r, column=c).value for c in range(1, limit + 1)])
            return rows

        for marker, limit in (("一、财务指标", 19), ("二、持仓个股财报摘要", 9)):
            merged = dump(ws, marker, limit)
            self.assertTrue(merged, f"合并页签缺少 {marker}")
            standalone = dump(ws_ind if marker.startswith("一") else ws_dig, marker, limit)
            for i, (m_row, s_row) in enumerate(zip(merged, standalone)):
                self.assertEqual(m_row, s_row, f"{marker} 第 {i} 行与独立写入不一致")
