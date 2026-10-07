"""Sheet 工厂 — 页签创建 + 可见性判定。

职责：根据注册表配置和运行时标志创建可见页签。
"""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Font

from src.python.core.registry import LLM_MODULE_GATED_SECTIONS as _LLM_MODULE_GATED_SECTIONS

# ── 页签分组配色（tabColor） ──────────────────────────────────
# 键 = core/registry.py::NAV_GROUPS 的组 key，值 = 6 位 hex 页签色。
# HTML 目录折叠分组（report/html_writer_nav.py）与本表同源于注册表条目
# nav_group 字段——同章同组必同色；键集与 NAV_GROUPS 双向一致由测试锁定
# （无孤儿组、无未登记组）。
_GROUP_TAB_COLORS: dict[str, str] = {
    "basic": "4472C4",  # 基础信息 — 蓝
    "fund_deep": "70AD47",  # 基金深度分析 — 绿
    "action": "ED7D31",  # 行动建议 — 橙
    "history": "FFC000",  # 历史 — 金
    "llm": "7030A0",  # LLM — 紫
    "appendix": "A5A5A5",  # 附录 — 灰
}


def apply_tab_color(ws: Any, section: dict) -> None:
    """按章节条目的 nav_group 给页签上色（未登记分组不上色，保持默认）。"""
    color = _GROUP_TAB_COLORS.get(section.get("nav_group", ""))
    if color:
        ws.sheet_properties.tabColor = color


def should_create_sheet(section: dict, data_availability: dict[str, bool] | None = None) -> bool:
    """纯 data 层：按注册表的 data_flag 判断模块数据是否就绪。

    无 data_flag 的模块（always、history）始终创建；data_flag
    未出现在 data_availability 中时视为已就绪（如基金深度分析的
    数据在页签创建后才写入，由下游函数自行兜底）。

    ``data_flag_any``（可选）为**多契约 OR** 口径，供章节合并后的条目使用：
    任一契约就绪即创建；未登记视为**未就绪**（悲观——否则合并条目在契约缺省时
    会创建空页签；单契约路径的乐观口径保持不变）。
    """
    avail = data_availability or {}
    flag_any = section.get("data_flag_any")
    if flag_any:
        return any(avail.get(name, False) for name in flag_any)
    flag_name = section.get("data_flag")
    if not flag_name:
        return True
    return avail.get(flag_name, True)


def build_data_availability(
    *,
    include_news: bool = False,
    include_llm: bool = False,
    enable_fund_deep_analysis: bool = False,
    financial_report_digest_data: dict | None = None,
    financial_indicator_data: dict | None = None,
    position_relationship_data: dict | None = None,
    holding_change_data: dict | None = None,
    schedule_replay_data: dict | None = None,
) -> dict[str, bool]:
    """构造 data 层可用性字典（章节可见性的单一事实来源）。

    注册表的 ``data_flag`` / ``data_flag_any`` 查此字典判定章节是否创建；
    合并章的契约 flag 口径集中在此处，避免各调用点（生成器、一致性测试镜像）
    各写一份而漂移：

      - 财报摘要 / 财务指标：契约非 None 即就绪（None = 对应功能开关关闭）
      - 持仓变动复盘：契约非 None 即就绪（None = 实验开关 `holding_change_review`
        关闭，页签不创建）
      - 调仓纪律回放：契约非 None 即就绪（None = 实验开关 `rebalance_schedule_replay`
        关闭/未注入，页签不创建；available=False 时页签写占位）
      - 持仓结构与集中度（合并章，两契约 OR）：基金深度分析开启时，重合度与集中度
        均由下游计算（数据不足时区块各自写占位）→ 两契约视为就绪；关闭时由 board 层隐藏
    """
    availability: dict[str, bool] = {}
    if include_news:
        availability["news_data_available"] = True
    if include_llm:
        availability["llm_data_available"] = True
    availability["financial_report_digest_data"] = financial_report_digest_data is not None
    availability["financial_indicator_data"] = financial_indicator_data is not None
    # 实验章：恒显式写入（True/False），避免 should_create_sheet 乐观缺省 True
    # 与 HTML 端 data_flags 悲观缺省 False 在键缺席时两端可见性不一致
    availability["holding_change_data"] = holding_change_data is not None
    availability["schedule_replay_data"] = schedule_replay_data is not None
    availability["position_relationship_data"] = enable_fund_deep_analysis or position_relationship_data is not None
    availability["concentration_data"] = enable_fund_deep_analysis
    return availability


def create_sheets(
    wb: Any,
    section_order: list[dict],
    enable_fund_deep_analysis: bool = False,
    enable_news: bool = True,  # board 层
    enable_history: bool = True,  # board 层
    enable_portfolio_evolution: bool = True,  # board 层：组合演进
    enable_fundamental_snapshot: bool = False,  # board 层：持仓基本面章（两功能开关任一开启）
    enable_action: bool = False,  # board 层：行动建议（config 默认开）
    enable_llm: bool = True,  # board 层
    data_availability: dict[str, bool] | None = None,  # data 层
    llm_module_disabled: dict[str, bool] | None = None,  # 章级：enabled_llm 模块禁用的 LLM 分析章不创建
) -> dict[str, Any]:
    """按配置顺序创建所有可见页签，返回 {key: ws} 字典。

    两层可见性模型：
      board 层：用户配置的章节开关（enable_xxx）
      data 层：运行时数据可用性标志（data_availability dict）

    Args:
        wb: openpyxl Workbook
        section_order: 注册表模块顺序列表
        enable_fund_deep_analysis: board 层 — 基金深度分析是否开启（配置驱动）
        enable_news: board 层 — 市场新闻是否开启（配置驱动）
        enable_history: board 层 — 历史走势章节是否开启
        enable_action: board 层 — 行动建议章节是否开启（config 默认开）
        enable_llm: board 层 — LLM 分析章节是否开启
        data_availability: data 层 — 各模块 data_flag 的就绪状态
        llm_module_disabled: 章级 — enabled_llm 逐模块禁用时，对应 LLM 分析章
            （core.registry.LLM_MODULE_GATED_SECTIONS）不创建页签（与 HTML 端同口径）
    """
    # 内联 board_flags dict（与 HTML 端结构一致，行为一致性由集成测试保证）
    board_flags = {
        "always": True,
        "fund_deep_analysis": enable_fund_deep_analysis,
        "news": enable_news,  # ← 配置驱动的 board 层值
        "history": enable_history,
        "evolution": enable_portfolio_evolution,
        # 持仓变动复盘：实验章无 board 层开关（恒 True），可见性由 data 层
        # data_flag（holding_change_data，seam 注入）控制
        "holding_change": True,
        # 调仓纪律回放：同持仓变动复盘（实验章无 board 层开关，data 层控制）
        "schedule_replay": True,
        "fundamental_snapshot": enable_fundamental_snapshot,
        "action": enable_action,
        "llm": enable_llm,
    }

    # should_create_sheet 查 data_availability dict
    sheets: dict[str, Any] = {}
    llm_usage_sec: dict | None = None
    visible_count = 0
    _data_avail = data_availability or {}

    for sec in section_order:
        # 第 1 层：board 层预过滤
        if not board_flags.get(sec.get("type", ""), True):
            continue

        # 第 1.5 层：章级（enabled_llm.<key> 禁用 → LLM 分析章整章隐藏，与 HTML 端同口径）
        if (
            llm_module_disabled
            and sec["key"] in _LLM_MODULE_GATED_SECTIONS
            and llm_module_disabled.get(sec["key"], False)
        ):
            continue

        # 第 2 层：data 层判断（查注册表的 data_flag）
        if not should_create_sheet(sec, _data_avail):
            continue

        # llm_usage 强制末位，先标记稍后创建
        if sec["key"] == "llm_usage":
            llm_usage_sec = sec
            continue

        visible_count += 1
        ws = wb.create_sheet()
        ws.title = f"{visible_count}.{sec['name']}"
        apply_tab_color(ws, sec)
        sheets[sec["key"]] = ws

    # llm_usage 始终在最后
    if llm_usage_sec is not None:
        visible_count += 1
        ws = wb.create_sheet()
        ws.title = f"{visible_count}.{llm_usage_sec['name']}"
        apply_tab_color(ws, llm_usage_sec)
        sheets["llm_usage"] = ws

    return sheets


def stamp_back_to_summary(sheets: dict[str, Any]) -> None:
    """给除汇总页外的每个页签标题行尾追加「↩ 返回汇总」内部超链接。

    写在标题行合并区之外（最后一列的下一列），不占新行、不移位后续
    内容与冻结窗格；汇总页本身为导航宿主不打标。须在各页签内容写完
    后调用（取 max_column 定位行尾）。
    """
    summary_ws = sheets.get("summary")
    if summary_ws is None:
        return
    summary_title = summary_ws.title
    for key, ws in sheets.items():
        if ws is None or key == "summary":
            continue
        cell = ws.cell(
            row=1,
            column=ws.max_column + 1,
            value=f'=HYPERLINK("#\'{summary_title}\'!A1","↩ 返回汇总")',
        )
        cell.font = Font(color="0563C1", underline="single")
