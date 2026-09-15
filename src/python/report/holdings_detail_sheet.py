"""持仓明细与分类 — Excel 写入层。

职责：把「市值核算明细」与「持仓分类汇总」两张表写入同一个 Excel 工作表，
分两个区块呈现：

  区块① 市值核算明细（按账户分组 + 分账户小计 + 总计，15/16 列）
  区块② 持仓分类汇总（资产属性 → 投资分类 → 明细 + 小计 + 总计，10/12 列）

仅依赖领域层类型与分类函数，不含行情获取/计算逻辑。

依赖方向：
  market_value.py  ← 本模块（DetailRow 类型）
  category.py      ← 本模块（分类/股息/数据状态领域函数）
  excel_generator.py / excel_content_sheets.py → 本模块（编排器调用）

分层边界：章节层（本模块）只做 Excel 呈现；分类领域函数
（`_categorize_holding` / `_tier_label` / `build_category_data_status` /
`calc_yield_text`）归 `category.py`，市值领域计算（`DetailRow` /
`classify_holdings` / `_compute_detail_row`）归 `market_value.py`，
本模块不重复实现。
"""

from __future__ import annotations

import logging
from typing import Any

from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from src.python.core.code_utils import is_qdii_extended
from src.python.core.models import Holding
from src.python.core.num_utils import finite_or
from src.python.core.registry import get_report_sheet_name
from src.python.report.category import (
    _categorize_holding,
    _load_dividend_data,
    build_category_data_status,
    calc_yield_text,
)
from src.python.report.excel_writer import (
    _write_data_status_foot,
    auto_width,
    freeze_header,
    write_data_row,
    write_header_row,
    write_subtotal_row,
    write_title_row,
    write_total_row,
)
from src.python.report.market_value import DetailRow
from src.python.report.styles import BLUE_FONT, FMT_MONEY, FMT_PERCENT, FMT_PRICE, FMT_SHARES, profit_font

logger = logging.getLogger("invest")

__all__ = [
    "write_holdings_detail_sheet",
    "_weighted_avg_cost",
    "_detail_to_row_values",
    "_apply_profit_colors",
    "_apply_price_type_colors",
]

# ── 区块① 市值核算明细 ──────────────────────────────────────
# 15 列表头（由 sheet 层维护）
_MV_HEADERS = [
    "账户",
    "名称",
    "代码",
    "最新价",
    "净值日期",
    "昨日价",
    "取价方式",
    "溢价率",
    "份额",
    "市值",
    "成本",
    "盈亏",
    "收益率",
    "本日盈亏",
    "取价渠道",
]
_MV_NCOLS = len(_MV_HEADERS)
# 成本流水子列（功能开关 `cost_lots` 开启时追加，默认关不渲染）
_MV_EXTRA_HEADERS = ["资金加权成本"]
_MV_NCOLS_WITH_FLOW = _MV_NCOLS + len(_MV_EXTRA_HEADERS)

_PRICE_TYPE_COL = 7
_NAME_COL = 2

# ── 区块② 持仓分类汇总 ──────────────────────────────────────
_CAT_NCOLS = 10
_CAT_HEADERS = [
    "资产属性",
    "投资分类",
    "名称",
    "代码",
    "市值",
    "成本",
    "盈亏",
    "收益率",
    "本日盈亏",
    "年均股息率",
]
# 成本流水子列（功能开关 `cost_lots` 开启时追加，默认关不渲染）
_CAT_EXTRA_HEADERS = ["成本分档", "分红累计"]
_CAT_NCOLS_WITH_FLOW = _CAT_NCOLS + len(_CAT_EXTRA_HEADERS)

# 区块小节标题
_BLOCK_TITLE_MARKET_VALUE = "一、市值核算明细"
_BLOCK_TITLE_CATEGORY = "二、持仓分类汇总"


def _weighted_avg_cost(buckets: dict | None) -> float | None:
    """资金加权成本（批次成本价按份额加权，含费用摊薄）。

    Args:
        buckets: cost_tiers.per_code 中某代码的分档桶 dict，缺码时为 None

    Returns:
        加权成本价浮点（无批次数据返回 None，渲染层写空）
    """
    if not isinstance(buckets, dict):
        return None
    shares = 0.0
    cost = 0.0
    for bucket in ("low", "high", "unpriced"):
        b = buckets.get(bucket) or {}
        shares += finite_or(b.get("shares", 0.0))
        cost += finite_or(b.get("cost", 0.0))
    if shares <= 0:
        return None
    return cost / shares


def _detail_to_row_values(d: DetailRow) -> list[Any]:
    """将 DetailRow 转为 Excel 行值列表。"""
    return [
        d.account,
        d.name,
        d.code,
        d.price,
        d.nav_date,
        d.yesterday_close,
        d.price_type,
        d.premium,
        d.shares,
        d.market_value,
        d.cost,
        d.profit,
        d.profit_rate,
        d.today_profit,
        d.source,
    ]


def _mv_num_formats(has_flow: bool = False) -> list[str | None]:
    """区块①每列的 Excel 数字格式。"""
    fmt = [
        "",  # 1  账户
        "",  # 2  名称
        "",  # 3  代码
        FMT_PRICE,  # 4  最新价
        "",  # 5  净值日期
        FMT_PRICE,  # 6  昨日价
        "",  # 7  取价方式
        "",  # 8  溢价率
        FMT_SHARES,  # 9  份额
        FMT_MONEY,  # 10 市值
        FMT_MONEY,  # 11 成本
        FMT_MONEY,  # 12 盈亏
        FMT_PERCENT,  # 13 收益率
        FMT_MONEY,  # 14 本日盈亏
        "",  # 15 取价渠道
    ]
    if has_flow:
        fmt += [FMT_PRICE]  # 16 资金加权成本（批次成本价按份额加权）
    return fmt


def _cat_num_formats(has_flow: bool = False) -> list[str | None]:
    """区块②每列的 Excel 数字格式。"""
    fmt = [
        "",  # 1  资产属性
        "",  # 2  投资分类
        "",  # 3  名称
        "",  # 4  代码
        FMT_MONEY,  # 5  市值
        FMT_MONEY,  # 6  成本
        FMT_MONEY,  # 7  盈亏
        FMT_PERCENT,  # 8  收益率
        FMT_MONEY,  # 9  本日盈亏
        "",  # 10 年均股息率（字符串格式）
    ]
    if has_flow:
        fmt += ["", FMT_MONEY]  # 11 成本分档（文本）、12 分红累计（金额）
    return fmt


def _apply_profit_colors(ws, start_row: int, end_row: int, profit_col: int, rate_col: int, today_col: int) -> None:
    """对区块①盈亏列（12）、收益率列（13）、本日盈亏列（14）着色。"""
    for r in range(start_row, end_row + 1):
        for col in (profit_col, today_col):
            cell = ws.cell(row=r, column=col)
            if isinstance(cell.value, (int, float)):
                cell.font = profit_font(cell.value)
        rate_cell = ws.cell(row=r, column=rate_col)
        if isinstance(rate_cell.value, float):
            rate_cell.font = profit_font(rate_cell.value)


def _apply_category_profit_colors(ws, start_row: int, end_row: int) -> None:
    """对区块②盈亏列 (7)、收益率列 (8)、本日盈亏列 (9) 着色。"""
    for r in range(start_row, end_row + 1):
        for col in (7, 9):
            cell = ws.cell(row=r, column=col)
            if isinstance(cell.value, (int, float)):
                cell.font = profit_font(cell.value)
        rate_cell = ws.cell(row=r, column=8)
        if isinstance(rate_cell.value, (int, float)):
            rate_cell.font = profit_font(rate_cell.value)


def _apply_price_type_colors(ws, start_row: int, end_row: int) -> None:
    """对取价方式列（第 7 列）着色：蓝色代表价格来源可靠/时效性高。"""
    for r in range(start_row, end_row + 1):
        cell = ws.cell(row=r, column=_PRICE_TYPE_COL)
        val = str(cell.value) if cell.value else ""
        if val in ("场内收盘价(T)", "场内午市收盘(T)", "官方净值(T)"):
            cell.font = BLUE_FONT
        elif val == "官方净值(T-1)":
            name_cell = ws.cell(row=r, column=_NAME_COL)
            name = str(name_cell.value) if name_cell.value else ""
            if is_qdii_extended(name):
                cell.font = BLUE_FONT


def _write_account_groupings(
    ws,
    details: list[DetailRow],
    data_start: int,
    fund_flow_data: dict | None = None,
) -> tuple[float, float, float, float, int]:
    """按账户分组写入明细行和小计，返回汇总数据及最终行号。

    fund_flow_data 非 None 时明细行追加「资金加权成本」列（数据契约）。

    Returns:
        (grand_mv, grand_cost, grand_profit, grand_today, final_row)
    """
    has_flow = fund_flow_data is not None
    ncols = _MV_NCOLS_WITH_FLOW if has_flow else _MV_NCOLS
    cost_map = (fund_flow_data or {}).get("cost_tiers", {}).get("per_code", {})

    accounts: dict[str, list[DetailRow]] = {}
    for d in details:
        accounts.setdefault(d.account, []).append(d)

    row = data_start
    grand_mv = grand_cost = grand_profit = grand_today = 0.0

    for acc_name, acc_details in accounts.items():
        for d in acc_details:
            vals = _detail_to_row_values(d)
            if has_flow:
                vals += [_weighted_avg_cost(cost_map.get(d.code))]
            write_data_row(ws, row, vals, _mv_num_formats(has_flow))
            row += 1

        acc_mv = sum(d.market_value for d in acc_details)
        acc_cost = sum(d.cost for d in acc_details)
        acc_profit = sum(d.profit for d in acc_details)
        acc_today = sum(d.today_profit for d in acc_details)
        acc_rate = acc_profit / acc_cost if acc_cost > 0 else 0.0

        subtotal_vals = [
            f"{acc_name} 小计",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            sum(d.shares for d in acc_details),
            acc_mv,
            acc_cost,
            acc_profit,
            acc_rate,
            acc_today,
            "",
        ]
        if has_flow:
            subtotal_vals += [""]  # 资金加权成本小计列留空（批次口径不跨账户聚合）
        write_subtotal_row(ws, row, f"{acc_name} 小计", subtotal_vals[1:], ncols, _mv_num_formats(has_flow))
        row += 1

        grand_mv += acc_mv
        grand_cost += acc_cost
        grand_profit += acc_profit
        grand_today += acc_today

    return grand_mv, grand_cost, grand_profit, grand_today, row


def _write_market_value_block(
    ws: Worksheet,
    details: list[DetailRow] | None,
    fund_flow_data: dict | None,
    start_row: int,
) -> tuple[float, float, float, float, list[DetailRow], int]:
    """写入区块①「市值核算明细」，返回 (汇总数据, 明细行, 下一可用行号)。"""
    has_flow = fund_flow_data is not None
    ncols = _MV_NCOLS_WITH_FLOW if has_flow else _MV_NCOLS
    headers = _MV_HEADERS + _MV_EXTRA_HEADERS if has_flow else _MV_HEADERS
    _details = details or []

    row = write_title_row(ws, start_row, _BLOCK_TITLE_MARKET_VALUE, ncols)
    row = write_header_row(ws, row, headers)
    data_start = row

    # 若所有行情数据全零，写一行醒目提示
    _all_zero = all(d.price == 0 for d in _details) if _details else False
    if _all_zero:
        _WARN_FONT = Font(size=10, bold=True, color="CC0000")
        ws.cell(row=row, column=1).font = _WARN_FONT
        ws.cell(row=row, column=1, value="⚠ 行情数据全部不可用（非交易时段/网络异常），以下市值/盈亏均为占位 —")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        row += 1

    # 按账户分组写入明细 + 小计
    grand_mv, grand_cost, grand_profit, grand_today, row = _write_account_groupings(
        ws, _details, data_start, fund_flow_data
    )

    # 总计
    grand_rate = grand_profit / grand_cost if grand_cost > 0 else 0.0
    total_vals = [
        "总计",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        sum(d.shares for d in _details) if _details else 0,
        grand_mv,
        grand_cost,
        grand_profit,
        grand_rate,
        grand_today,
        "",
    ]
    if has_flow:
        total_vals += [""]  # 资金加权成本总计列留空（批次口径不跨账户聚合）
    write_total_row(ws, row, "总计", total_vals[1:], ncols, _mv_num_formats(has_flow))

    # 对盈亏列着色
    _apply_profit_colors(ws, data_start, row, profit_col=12, rate_col=13, today_col=14)

    # 对取价方式列着色
    _apply_price_type_colors(ws, data_start, row)

    logger.info(
        "%s 区块①写入完成，共 %d 个账户，%d 条持仓",
        _BLOCK_TITLE_MARKET_VALUE,
        len(set(d.account for d in _details)),
        len(_details),
    )
    return grand_mv, grand_cost, grand_profit, grand_today, _details, row + 1


def _write_category_group(
    ws: Worksheet,
    row: int,
    group: list[Holding],
    prop: str,
    sub: str,
    detail_map: dict,
    dividend_data: dict,
    fund_flow_data: dict | None = None,
) -> tuple[int, float, float, float, float, float]:
    """写入一个分类分组的明细行和小计。

    fund_flow_data 非 None 时追加「成本分档」「分红累计」子列（数据契约）。

    Returns:
        (next_row, mv, cost, profit, today, div_sum)
    """
    has_flow = fund_flow_data is not None
    tier_map = (fund_flow_data or {}).get("cost_tiers", {}).get("per_code", {})
    div_map = (fund_flow_data or {}).get("dividends", {}).get("per_code", {})
    div_sum = 0.0

    for h in group:
        d = detail_map.get(h.code)
        if d:
            vals = [
                prop,
                sub,
                h.name,
                h.code,
                d.market_value,
                d.cost,
                d.profit,
                d.profit_rate,
                d.today_profit,
                calc_yield_text(h.code, d, dividend_data),
            ]
        else:
            vals = [prop, sub, h.name, h.code, 0.0, 0.0, 0.0, 0.0, 0.0, "--"]
        if has_flow:
            from src.python.report.category import _tier_label

            vals += [_tier_label(tier_map.get(h.code)), div_map.get(h.code, 0.0)]
            div_sum += div_map.get(h.code, 0.0)
        write_data_row(ws, row, vals, _cat_num_formats(has_flow))
        row += 1

    sub_mv = sum(detail_map.get(h.code, DetailRow()).market_value for h in group if h.code in detail_map)
    sub_cost = sum(detail_map.get(h.code, DetailRow()).cost for h in group if h.code in detail_map)
    sub_profit = sum(detail_map.get(h.code, DetailRow()).profit for h in group if h.code in detail_map)
    sub_today = sum(detail_map.get(h.code, DetailRow()).today_profit for h in group if h.code in detail_map)
    sub_rate = sub_profit / sub_cost if sub_cost > 0 else 0.0

    subtotal_vals = ["", "", len(group), sub_mv, sub_cost, sub_profit, sub_rate, sub_today, "--"]
    if has_flow:
        subtotal_vals += ["", div_sum]
    ncols = _CAT_NCOLS_WITH_FLOW if has_flow else _CAT_NCOLS
    write_subtotal_row(ws, row, f"{prop} - {sub} 小计", subtotal_vals, ncols, _cat_num_formats(has_flow))
    return row + 1, sub_mv, sub_cost, sub_profit, sub_today, div_sum


def _write_category_block(
    ws: Worksheet,
    holdings: list[Holding],
    details: list[DetailRow],
    fund_flow_data: dict | None,
    start_row: int,
) -> int:
    """写入区块②「持仓分类汇总」，返回下一可用行号。

    分类层级：资产属性 → 投资分类 → 持仓明细 → 小计 → 总计。
    每行含市值、成本、盈亏、收益率、本日盈亏、年均股息率。
    """
    has_flow = fund_flow_data is not None
    ncols = _CAT_NCOLS_WITH_FLOW if has_flow else _CAT_NCOLS
    headers = _CAT_HEADERS + _CAT_EXTRA_HEADERS if has_flow else _CAT_HEADERS

    detail_map: dict[str, DetailRow] = {d.code: d for d in details}

    cat_groups: dict[tuple[str, str], list[Holding]] = {}
    for h in holdings:
        prop, sub = _categorize_holding(h)
        cat_groups.setdefault((prop, sub), []).append(h)

    _PROP_ORDER = {"股票": 0, "基金": 1, "债券": 2, "现金": 3, "其他": 4}
    _SUB_ORDER = {"A股": 0, "QDII": 1, "主动": 2, "被动": 3, "指数": 4, "混合": 5, "纯债": 6, "货币": 7, "其他": 8}
    sorted_groups = sorted(
        cat_groups.items(),
        key=lambda x: (_PROP_ORDER.get(x[0][0], 99), _SUB_ORDER.get(x[0][1], 99)),
    )

    row = write_title_row(ws, start_row, _BLOCK_TITLE_CATEGORY, ncols)
    row = write_header_row(ws, row, headers)
    data_start = row

    # 若所有行情数据全零，写一行醒目提示
    _all_zero = all(d.market_value == 0 for d in details)
    if _all_zero and details:
        cell = ws.cell(row=row, column=1, value="⚠ 行情数据全部不可用（非交易时段/网络异常），以下市值/盈亏均为占位 —")
        cell.font = Font(size=10, bold=True, color="CC0000")
        row += 1

    dividend_data, dividend_success = _load_dividend_data(holdings)
    grand_mv = grand_cost = grand_profit = grand_today = grand_div = 0.0

    for (prop, sub), group in sorted_groups:
        row, smv, scost, sprofit, stoday, sdiv = _write_category_group(
            ws,
            row,
            group,
            prop,
            sub,
            detail_map,
            dividend_data,
            fund_flow_data,
        )
        grand_mv += smv
        grand_cost += scost
        grand_profit += sprofit
        grand_today += stoday
        grand_div += sdiv

    grand_rate = grand_profit / grand_cost if grand_cost > 0 else 0.0
    total_vals = ["", "", "-", grand_mv, grand_cost, grand_profit, grand_rate, grand_today, "--"]
    if has_flow:
        total_vals += ["", grand_div]
    write_total_row(ws, row, "总计", total_vals, ncols, _cat_num_formats(has_flow))

    _apply_category_profit_colors(ws, data_start, row)

    logger.info(
        "%s 区块②写入完成，共 %d 个分组，%d 条持仓",
        _BLOCK_TITLE_CATEGORY,
        len(sorted_groups),
        len(holdings),
    )

    data_status = build_category_data_status(dividend_success)
    _write_data_status_foot(ws, data_status, start_row=row + 1, max_cols=ncols)
    return row + 1


def write_holdings_detail_sheet(
    ws: Worksheet,
    holdings: list[Holding] | None = None,
    details: list[DetailRow] | None = None,
    fund_flow_data: dict | None = None,
) -> tuple[float, float, float, float, list[DetailRow]]:
    """写入「持仓明细与分类」工作表（区块①市值明细 + 区块②分类汇总）。

    Args:
        ws: 目标工作表
        holdings: 原始持仓列表（区块②分类分组用；None 视为空）
        details: 预计算明细行（必须传入，由编排器预计算）
        fund_flow_data: 成本流水数据契约（非 None 时两区块各追加子列；
            None 时保持既有 15 列 / 10 列输出）

    Returns:
        (总市值, 总成本, 总盈亏, 本日总盈亏, 明细行列表)
    """
    _details = details or []
    _holdings = holdings or []

    row = write_title_row(ws, 1, get_report_sheet_name("holdings_detail"), _MV_NCOLS_WITH_FLOW)

    grand_mv, grand_cost, grand_profit, grand_today, _details, next_row = _write_market_value_block(
        ws, _details, fund_flow_data, row
    )

    _write_category_block(ws, _holdings, _details, fund_flow_data, next_row + 1)

    freeze_header(ws, 3)
    auto_width(ws)

    logger.info(
        "%s写入完成，共 %d 个账户，%d 条持仓",
        get_report_sheet_name("holdings_detail"),
        len(set(d.account for d in _details)),
        len(_details),
    )
    return grand_mv, grand_cost, grand_profit, grand_today, _details
